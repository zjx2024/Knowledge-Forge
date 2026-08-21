"""Embedded document image extraction and knowledge enrichment."""

from __future__ import annotations

import hashlib
import io
import json
import logging
import threading
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageFile, ImageOps

from .chunking import (
    DocumentBlock,
    ParsedDocument,
    _heading_info,
    _heading_path,
    _slug,
    sanitize_parsed_document,
)
from .config import Config
from .ollama import OllamaClient
from .unicode_safety import sanitize_unicode_data


LOGGER = logging.getLogger(__name__)
ImageFile.LOAD_TRUNCATED_IMAGES = True
IMAGE_ANALYSIS_CACHE_VERSION = 3


IMAGE_ANALYSIS_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "description": {"type": "string"},
        "ocr_text": {"type": "string"},
        "facts": {"type": "array", "items": {"type": "string"}},
        "relationships": {"type": "array", "items": {"type": "string"}},
        "knowledge_relevance": {
            "type": "string",
            "enum": ["substantive", "context_only", "decorative"],
        },
        "relevance_reason": {"type": "string"},
    },
    "required": [
        "title", "description", "ocr_text", "facts", "relationships",
        "knowledge_relevance", "relevance_reason",
    ],
}


@dataclass
class EmbeddedImage:
    data: bytes
    content_hash: str
    width: int
    height: int
    sources: list[dict[str, Any]] = field(default_factory=list)
    local_path: Path | None = None


class DocumentMediaService:
    """Turn embedded PDF/DOCX images into provenance-rich text blocks."""

    _ocr_lock = threading.Lock()
    _ocr_engine: Any = None

    def __init__(self, ollama: OllamaClient):
        self.ollama = ollama

    def enrich(
        self,
        path: Path,
        parsed: ParsedDocument,
        project: dict[str, Any],
        media_dir: Path,
        *,
        progress: Callable[[int, int, str], None] | None = None,
    ) -> ParsedDocument:
        if path.suffix.lower() not in {".pdf", ".docx"}:
            return sanitize_parsed_document(parsed)
        media_dir.mkdir(parents=True, exist_ok=True)
        raw = self._extract(path, parsed)
        assets = self._materialize(raw, media_dir)
        image_blocks: list[DocumentBlock] = []
        failures = 0
        ocr_chars = 0
        for index, asset in enumerate(assets, start=1):
            source = asset.sources[0] if asset.sources else {}
            try:
                analysis = self._analyze(asset, project)
                text = self._analysis_text(analysis, source)
                ocr_chars += len(str(analysis.get("ocr_text") or ""))
                if text:
                    order = len(parsed.blocks) + len(image_blocks) + 1
                    metadata = {
                        "asset_type": "document_image",
                        "image_hash": asset.content_hash,
                        "local_path": str(asset.local_path),
                        "url": Config.file_url(asset.local_path),
                        "width": asset.width,
                        "height": asset.height,
                        "sources": asset.sources,
                        "analysis_model": project.get("llm_model") or Config.default_llm_model,
                        "title": str(analysis.get("title") or ""),
                        "description": str(analysis.get("description") or ""),
                        "knowledge_relevance": str(
                            analysis.get("knowledge_relevance") or "context_only"
                        ),
                        "relevance_reason": str(
                            analysis.get("relevance_reason") or ""
                        ),
                        "graph_extraction_enabled": (
                            str(analysis.get("knowledge_relevance") or "")
                            == "substantive"
                        ),
                    }
                    image_blocks.append(DocumentBlock(
                        key=f"b{order:05d}",
                        block_type="image",
                        text=text,
                        order=order,
                        # Keep every media analysis in its own retrieval region
                        # so a non-substantive image can be excluded from graph
                        # extraction without disabling adjacent body text.
                        section_key=(
                            f"{str(source.get('section_key') or 'media')}"
                            f"_image_{asset.content_hash[:12]}"
                        ),
                        section_title=str(source.get("section_title") or path.stem),
                        level=int(source.get("level") or 1),
                        heading_path=tuple(source.get("heading_path") or (path.stem,)),
                        page=int(source["page"]) if source.get("page") is not None else None,
                        metadata=metadata,
                    ))
            except Exception as exc:
                failures += 1
                LOGGER.warning("Image analysis failed for %s in %s: %s", asset.content_hash[:12], path.name, exc)
            if progress:
                progress(index, len(assets), f"图片内容 {index}/{len(assets)}：{self._source_label(source)}")

        quality = dict(parsed.quality)
        quality.update({
            "image_count": len(assets),
            "image_analyzed_count": len(image_blocks),
            "image_analysis_failures": failures,
            "image_ocr_chars": ocr_chars,
            "block_count": len(parsed.blocks) + len(image_blocks),
            "text_chars": int(parsed.quality.get("text_chars") or 0) + sum(len(item.text) for item in image_blocks),
        })
        warnings = list(quality.get("warnings") or [])
        if failures:
            warnings.append(f"{failures} 张图片未能完成视觉理解")
        quality["warnings"] = warnings
        return sanitize_parsed_document(ParsedDocument(
            blocks=[*parsed.blocks, *image_blocks],
            parser_type=parsed.parser_type,
            parser_version=parsed.parser_version,
            quality=quality,
        ))

    def _extract(self, path: Path, parsed: ParsedDocument) -> list[tuple[bytes, dict[str, Any]]]:
        if path.suffix.lower() == ".docx":
            return self._extract_docx(path)
        return self._extract_pdf(path, parsed)

    @staticmethod
    def _paragraph_images(paragraph: Paragraph) -> Iterable[bytes]:
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship_id = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if not relationship_id:
                continue
            part = paragraph.part.related_parts.get(relationship_id)
            blob = getattr(part, "blob", None)
            if blob:
                yield bytes(blob)

    def _extract_docx(self, path: Path) -> list[tuple[bytes, dict[str, Any]]]:
        document = Document(path)
        values: list[tuple[bytes, dict[str, Any]]] = []
        raw_hashes: set[str] = set()
        title, level, section_index = path.stem, 1, 1
        section_key = _slug(title, section_index)
        stack = [title]

        def add_paragraph(paragraph: Paragraph, body_order: int, location: str) -> None:
            for image_index, data in enumerate(self._paragraph_images(paragraph), start=1):
                raw_hashes.add(hashlib.sha256(data).hexdigest())
                values.append((data, {
                    "source_kind": "docx_embedded",
                    "body_order": body_order,
                    "image_index": image_index,
                    "location": location,
                    "section_key": section_key,
                    "section_title": title,
                    "level": level,
                    "heading_path": tuple(stack),
                }))

        for body_order, item in enumerate(document.iter_inner_content(), start=1):
            if isinstance(item, Paragraph):
                is_heading, candidate_level, text = _heading_info(item)
                if is_heading:
                    title, level = text, candidate_level
                    section_index += 1
                    section_key = _slug(title, section_index)
                    _heading_path(stack, level, title)
                add_paragraph(item, body_order, "paragraph")
            elif isinstance(item, Table):
                seen_cells: set[int] = set()
                for row_index, row in enumerate(item.rows, start=1):
                    for cell_index, cell in enumerate(row.cells, start=1):
                        cell_id = id(cell._tc)
                        if cell_id in seen_cells:
                            continue
                        seen_cells.add(cell_id)
                        for paragraph in cell.paragraphs:
                            add_paragraph(paragraph, body_order, f"table:{row_index}:{cell_index}")

        # Headers, floating drawings and unusual placements may not appear in
        # iter_inner_content. Preserve them with document-level provenance.
        try:
            with zipfile.ZipFile(path) as archive:
                for name in sorted(value for value in archive.namelist() if value.startswith("word/media/")):
                    data = archive.read(name)
                    if hashlib.sha256(data).hexdigest() in raw_hashes:
                        continue
                    values.append((data, {
                        "source_kind": "docx_embedded",
                        "location": name,
                        "section_key": _slug(path.stem, 1),
                        "section_title": path.stem,
                        "level": 1,
                        "heading_path": (path.stem,),
                    }))
        except (OSError, zipfile.BadZipFile):
            pass
        return values

    def _extract_pdf(self, path: Path, parsed: ParsedDocument) -> list[tuple[bytes, dict[str, Any]]]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        page_sections = {
            int(block.page): block
            for block in parsed.blocks
            if block.page is not None and block.block_type != "heading"
        }
        values: list[tuple[bytes, dict[str, Any]]] = []
        render_pages: list[int] = []
        for page_index, page in enumerate(reader.pages, start=1):
            section = page_sections.get(page_index)
            source = {
                "source_kind": "pdf_embedded",
                "page": page_index,
                "section_key": section.section_key if section else f"media_p{page_index:04d}",
                "section_title": section.section_title if section else f"第 {page_index} 页图片",
                "level": section.level if section else 1,
                "heading_path": section.heading_path if section else (path.stem, f"第 {page_index} 页"),
            }
            page_image_count = 0
            try:
                for image_index, image in enumerate(page.images, start=1):
                    data = bytes(image.data)
                    values.append((data, {**source, "image_index": image_index,
                                           "location": str(getattr(image, "name", ""))}))
                    page_image_count += 1
            except Exception as exc:
                LOGGER.warning("PDF image extraction failed on page %s of %s: %s", page_index, path.name, exc)
            text = page.extract_text() or ""
            if page_image_count == 0 and len(text.strip()) <= Config.image_scan_text_threshold:
                render_pages.append(page_index)

        if render_pages:
            try:
                import pypdfium2 as pdfium

                pdf = pdfium.PdfDocument(str(path))
                try:
                    for page_index in render_pages:
                        page = pdf[page_index - 1]
                        rendered = page.render(scale=Config.image_pdf_render_scale).to_pil()
                        buffer = io.BytesIO()
                        rendered.save(buffer, format="PNG")
                        section = page_sections.get(page_index)
                        values.append((buffer.getvalue(), {
                            "source_kind": "pdf_scanned_page",
                            "page": page_index,
                            "image_index": 1,
                            "location": "rendered_page",
                            "section_key": section.section_key if section else f"media_p{page_index:04d}",
                            "section_title": section.section_title if section else f"第 {page_index} 页图片",
                            "level": section.level if section else 1,
                            "heading_path": section.heading_path if section else (path.stem, f"第 {page_index} 页"),
                        }))
                finally:
                    pdf.close()
            except Exception as exc:
                LOGGER.warning("Scanned PDF page rendering failed for %s: %s", path.name, exc)
        return values

    def _materialize(self, raw: list[tuple[bytes, dict[str, Any]]], media_dir: Path) -> list[EmbeddedImage]:
        media_dir.mkdir(parents=True, exist_ok=True)
        unique: dict[str, EmbeddedImage] = {}
        for data, source in raw:
            normalized = self._normalize(data)
            if not normalized:
                continue
            image_data, width, height, content_hash = normalized
            if content_hash in unique:
                unique[content_hash].sources.append(source)
                continue
            target = media_dir / f"{content_hash}.png"
            if not target.exists():
                target.write_bytes(image_data)
            unique[content_hash] = EmbeddedImage(
                data=image_data,
                content_hash=content_hash,
                width=width,
                height=height,
                sources=[source],
                local_path=target,
            )
        return list(unique.values())

    @staticmethod
    def _normalize(data: bytes) -> tuple[bytes, int, int, str] | None:
        try:
            with Image.open(io.BytesIO(data)) as loaded:
                image = ImageOps.exif_transpose(loaded)
                image.load()
                width, height = image.size
                if (width < Config.image_min_width or height < Config.image_min_height
                        or width * height < Config.image_min_pixels):
                    return None
                if image.mode in {"RGBA", "LA"} or (image.mode == "P" and "transparency" in image.info):
                    rgba = image.convert("RGBA")
                    background = Image.new("RGBA", rgba.size, "white")
                    background.alpha_composite(rgba)
                    image = background.convert("RGB")
                else:
                    image = image.convert("RGB")
                image.thumbnail((Config.image_max_dimension, Config.image_max_dimension), Image.Resampling.LANCZOS)
                width, height = image.size
                buffer = io.BytesIO()
                image.save(buffer, format="PNG", optimize=True)
                normalized = buffer.getvalue()
                return normalized, width, height, hashlib.sha256(normalized).hexdigest()
        except Exception:
            return None

    def _analyze(self, asset: EmbeddedImage, project: dict[str, Any]) -> dict[str, Any]:
        model = project.get("llm_model") or Config.default_llm_model
        cache_path = asset.local_path.with_suffix(".analysis.json")
        if cache_path.exists():
            try:
                cached = json.loads(cache_path.read_text(encoding="utf-8"))
                if (
                    cached.get("model") == model
                    and cached.get("version") == IMAGE_ANALYSIS_CACHE_VERSION
                ):
                    return sanitize_unicode_data(
                        dict(cached.get("analysis") or {})
                    )
            except (OSError, json.JSONDecodeError):
                pass

        ocr_text = self._ocr(asset.data) if Config.image_ocr_enabled else ""
        analysis: dict[str, Any] = {
            "title": "",
            "description": "",
            "ocr_text": ocr_text,
            "facts": [],
            "relationships": [],
            # Without visual analysis, preserve the previous graph behavior;
            # OCR-only deployments cannot safely infer that an image is merely
            # decorative. The configured vision pass supplies the real role.
            "knowledge_relevance": "substantive",
            "relevance_reason": "visual relevance classification unavailable",
        }
        if Config.image_analysis_enabled:
            source = asset.sources[0] if asset.sources else {}
            visual = self.ollama.json_chat(
                model=model,
                system=(
                    "You extract knowledge from an image embedded in a source document. Describe only visible content. "
                    "Transcribe meaningful text, list explicit factual statements, and express visible relationships as "
                    "short 'A --relation--> B' strings. Classify knowledge_relevance as substantive only when the image "
                    "itself conveys domain facts useful to a knowledge graph, such as a diagram, chart, data table, map, "
                    "technical screenshot, or evidentiary photograph. Use context_only for portraits, author photos, "
                    "signatures, covers and logos; use decorative for styling or layout imagery. Non-substantive images "
                    "remain searchable/displayable but must not contribute graph facts. Relationships must express "
                    "domain semantics between meaningful objects or concepts. Keep axis names, legends, colours, marker "
                    "shapes and other visual encodings in the description or OCR text, but do not emit them as graph "
                    "relationships. Do not infer facts that are not "
                    "shown. Return schema-valid JSON only."
                ),
                prompt=(
                    f"Image source: {self._source_label(source)}\n"
                    f"Independent OCR text (may contain errors):\n{ocr_text[:12000]}\n\n"
                    "Analyze the attached document image. Ignore decorative styling and report its useful knowledge."
                ),
                schema=IMAGE_ANALYSIS_SCHEMA,
                temperature=0,
                context_budget_tokens=int(project.get("model_max_tokens") or 32768),
                max_output_tokens=Config.image_analysis_max_output_tokens,
                images=[asset.data],
            )
            analysis.update({key: visual.get(key, analysis[key]) for key in analysis})
            if ocr_text and not str(analysis.get("ocr_text") or "").strip():
                analysis["ocr_text"] = ocr_text

        analysis = sanitize_unicode_data(analysis)
        cache_path.write_text(json.dumps({
            "version": IMAGE_ANALYSIS_CACHE_VERSION,
            "model": model,
            "analysis": analysis,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        return analysis

    @classmethod
    def _ocr(cls, data: bytes) -> str:
        try:
            with cls._ocr_lock:
                if cls._ocr_engine is None:
                    from rapidocr import RapidOCR
                    cls._ocr_engine = RapidOCR()
                result = cls._ocr_engine(data)
            texts = getattr(result, "txts", None) or []
            return "\n".join(str(value).strip() for value in texts if str(value).strip())
        except Exception as exc:
            LOGGER.warning("Image OCR unavailable: %s", exc)
            return ""

    @staticmethod
    def _analysis_text(analysis: dict[str, Any], source: dict[str, Any]) -> str:
        lines = ["[文档图片]", f"图片来源: {DocumentMediaService._source_label(source)}"]
        fields = (("图片标题", analysis.get("title")),
                  ("图片描述", analysis.get("description")),
                  ("图片文字", analysis.get("ocr_text")))
        lines.extend(f"{label}: {str(value).strip()}" for label, value in fields if str(value or "").strip())
        facts = [str(value).strip() for value in analysis.get("facts") or [] if str(value).strip()]
        relations = [str(value).strip() for value in analysis.get("relationships") or [] if str(value).strip()]
        if facts:
            lines.append("图片中明确事实:\n" + "\n".join(f"- {value}" for value in facts))
        if relations:
            lines.append("图片中明确关系:\n" + "\n".join(f"- {value}" for value in relations))
        return "\n".join(lines).strip() if len(lines) > 2 else ""

    @staticmethod
    def _source_label(source: dict[str, Any]) -> str:
        if source.get("page") is not None:
            return f"PDF 第 {source['page']} 页"
        section = str(source.get("section_title") or "文档")
        location = str(source.get("location") or "内嵌图片")
        return f"DOCX 章节“{section}” ({location})"
