"""Deterministic document routing, extraction and hierarchy-aware chunking."""

from __future__ import annotations

import csv
import gc
import hashlib
import importlib.metadata
import importlib.util
import logging
import math
import re
import threading
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .config import Config
from .entity_publishability import bibliography_reference_spans
from .sentence_spans import split_source_sentences
from .unicode_safety import (
    sanitize_unicode_data,
    sanitize_unicode_text,
    unsafe_unicode_code_unit_count,
)


LOGGER = logging.getLogger(__name__)
_DOCLING_LOCK = threading.Lock()


PARSER_OPTIONS = {
    "auto": "自动选择",
    "native_docx": "Word 原生解析",
    "native_text": "文本/Markdown 解析",
    "plaintext_pdf": "PDF 文本层解析",
    "docling": "Docling 版面解析",
    "spreadsheet": "电子表格解析",
}
CHUNK_STRATEGY_OPTIONS = {
    "auto": "自动选择",
    "default": "兼容分块",
    "general": "通用分块",
    "paper": "论文分块",
    "table": "表格分块",
}

# These versions are part of the persisted document-revision fingerprint.
# Increment them whenever structure interpretation or chunk construction can
# change for identical source bytes and user settings.
DOCUMENT_STRUCTURE_PROFILE_VERSION = (
    "kf-document-structure-v8-docling-heading-sanity-all-routes"
)
DOCUMENT_CHUNK_PROFILE_VERSION = (
    "kf-document-chunks-v12-split-adjacent-article-preview"
)

_PARSER_SUFFIXES = {
    "native_docx": {".docx"},
    "native_text": {".md", ".markdown", ".txt"},
    "plaintext_pdf": {".pdf"},
    "docling": {".pdf", ".docx"},
    "spreadsheet": {".xlsx", ".xls", ".csv"},
}
_STRATEGY_SUFFIXES = {
    "default": {".docx", ".pdf", ".md", ".markdown", ".txt"},
    "general": set(Config.allowed_suffixes),
    "paper": {".docx", ".pdf", ".md", ".markdown", ".txt"},
    "table": {".docx", ".pdf", ".xlsx", ".xls", ".csv"},
}
_DOTTED_NUMBERED_HEADING = re.compile(
    r"^(\d+(?:\.\d+){1,7})(?:\.\s+|[、．]\s*|\s+)(.+)$"
)
_SPACED_DOTTED_NUMBERED_HEADING = re.compile(
    r"^(\d+(?:\s*\.\s*\d+){1,7})(?:\s+|[、．]\s*)(.+)$"
)
_SINGLE_NUMBERED_HEADING = re.compile(
    r"^(\d+)(?:[、．]\s*|\.(?!\d)\s*|\s+)(.+)$"
)
_TOC_DOT_LEADER = re.compile(r"(?:\.{4,}|…{2,})\s*\d*\s*$")
_TOC_ENTRY_DOT_LEADER = re.compile(r"(?:\.{4,}|…{2,})\s*\d+\s*$")
_SYNTHETIC_PDF_PAGE_TITLE = re.compile(r"^第\s*\d+\s*页$")
_DOCLING_DECIMAL_MEASUREMENT_HEADING = re.compile(
    r"^[+−-]?\d+[.,]\d{2,}\s*"
    r"(?:%|‰|ns|μs|us|ms|sec|secs|s|min|mins|h|hz|khz|mhz|ghz|"
    r"kb|mb|gb|tb|b|mm|cm|km|mg|kg|db|fps|w|kw|v|mv|a|ma)"
    r"(?:\s+[A-Za-z][A-Za-z0-9_.+-]{0,31})?\s*$",
    re.IGNORECASE,
)
_PAPER_MARKERS = re.compile(
    r"(?:摘要|关键词|引言|研究方法|实验(?:结果)?|结论|参考文献|abstract|keywords?|introduction|methodology|results?|conclusion|references)",
    re.IGNORECASE,
)
_REFERENCE_TITLE = re.compile(r"^(?:参考文献|references?|bibliography)\s*$", re.IGNORECASE)
_ABSTRACT_TITLE = re.compile(r"^(?:摘要|abstract)\s*$", re.IGNORECASE)
_ZH_ABSTRACT_MARKER = re.compile(
    r"(?im)^[ \t\u3000]*[【\[（(]?[ \t\u3000]*摘[ \t\u3000]*要"
    r"[ \t\u3000]*[】\]）)]?[ \t\u3000]*[:：]?[ \t\u3000]*"
)
_OCR_INLINE_SPACE = r"[ \t\u3000]*"
_EN_ABSTRACT_MARKER = re.compile(
    rf"(?im)^[ \t\u3000]*[【\[（(]?[ \t\u3000]*"
    rf"a{_OCR_INLINE_SPACE}b{_OCR_INLINE_SPACE}s{_OCR_INLINE_SPACE}t"
    rf"{_OCR_INLINE_SPACE}r{_OCR_INLINE_SPACE}a{_OCR_INLINE_SPACE}c"
    rf"{_OCR_INLINE_SPACE}t[ \t\u3000]*[】\]）)]?"
    rf"[ \t\u3000]*[:：]?[ \t\u3000]*"
)
_ZH_KEYWORD_MARKER = re.compile(
    r"(?im)^[ \t\u3000]*[【\[（(]?[ \t\u3000]*关[ \t\u3000]*键[ \t\u3000]*词"
    r"[ \t\u3000]*[】\]）)]?[ \t\u3000]*[:：]?[ \t\u3000]*"
)
_EN_KEYWORD_MARKER = re.compile(
    rf"(?im)^[ \t\u3000]*[【\[（(]?[ \t\u3000]*"
    rf"k{_OCR_INLINE_SPACE}e{_OCR_INLINE_SPACE}y{_OCR_INLINE_SPACE}w"
    rf"{_OCR_INLINE_SPACE}o{_OCR_INLINE_SPACE}r{_OCR_INLINE_SPACE}d"
    rf"{_OCR_INLINE_SPACE}s?[ \t\u3000]*[】\]）)]?"
    rf"[ \t\u3000]*[:：]?[ \t\u3000]*"
)
_AUTHOR_BIOGRAPHY_MARKER = re.compile(
    r"(?im)^[ \t\u3000]*(?:"
    r"作[ \t\u3000]*者[ \t\u3000]*(?:简[ \t\u3000]*介|介[ \t\u3000]*绍)"
    r"|author(?:s|'s)?[ \t]+(?:biograph(?:y|ies)|information)"
    r")[ \t\u3000]*[:：]?[ \t\u3000]*"
)
_DOI_MARKER = re.compile(r"(?i)\bdoi\s*[:：]\s*10\.\d{3,9}/")
_BILINGUAL_SUMMARY_TITLES: tuple[tuple[re.Pattern[str], str, str], ...] = (
    (re.compile(r"^(?:执行摘要)$"), "zh", "executive_summary"),
    (re.compile(r"^(?:executive[ \t]+summary)$", re.IGNORECASE), "en", "executive_summary"),
    (re.compile(r"^(?:内容提要|内容摘要)$"), "zh", "summary"),
    (re.compile(r"^(?:summary)$", re.IGNORECASE), "en", "summary"),
    (re.compile(r"^(?:概述|概览)$"), "zh", "overview"),
    (re.compile(r"^(?:overview)$", re.IGNORECASE), "en", "overview"),
)


@dataclass(frozen=True)
class DocumentBlock:
    key: str
    block_type: str
    text: str
    order: int
    section_key: str
    section_title: str
    level: int = 1
    heading_path: tuple[str, ...] = ()
    page: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class DocumentProfile:
    suffix: str
    page_count: int = 0
    sampled_pages: int = 0
    text_chars: int = 0
    has_text_layer: bool = True
    empty_page_ratio: float = 0.0
    replacement_ratio: float = 0.0
    heading_count: int = 0
    table_count: int = 0
    image_count: int = 0
    paper_score: int = 0
    layout_complexity: str = "simple"
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ProcessingPlan:
    requested_parser: str
    requested_strategy: str
    effective_parser: str
    effective_strategy: str
    routing_reason: str
    profile: DocumentProfile


@dataclass(frozen=True)
class ParsedDocument:
    blocks: list[DocumentBlock]
    parser_type: str
    parser_version: str
    quality: dict[str, Any]


@dataclass(frozen=True)
class Section:
    key: str
    title: str
    level: int
    content: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class TextChunk:
    key: str
    section_key: str
    title: str
    part_index: int
    part_total: int
    content: str
    content_hash: str
    token_count: int = 0
    parent_key: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class _BilingualSummarySpan:
    section_index: int
    start: int
    end: int
    language: str
    abstract_text: str
    kind: str = "abstract"


@dataclass(frozen=True)
class ProcessingResult:
    plan: ProcessingPlan
    parsed: ParsedDocument
    sections: list[Section]
    chunks: list[TextChunk]

    def payload(self) -> dict[str, Any]:
        return {
            "plan": {**asdict(self.plan), "profile": asdict(self.plan.profile)},
            "quality": self.parsed.quality,
            "blocks": [asdict(item) for item in self.parsed.blocks],
            "sections": [asdict(item) for item in self.sections],
            "chunks": [asdict(item) for item in self.chunks],
        }


def sanitize_parsed_document(parsed: ParsedDocument) -> ParsedDocument:
    """Repair parser/OCR text before it reaches UTF-8 or PostgreSQL."""

    unsafe_count = unsafe_unicode_code_unit_count({
        "blocks": [asdict(item) for item in parsed.blocks],
        "parser_type": parsed.parser_type,
        "parser_version": parsed.parser_version,
        "quality": parsed.quality,
    })
    if not unsafe_count:
        return parsed
    blocks = [
        DocumentBlock(
            key=sanitize_unicode_text(item.key),
            block_type=sanitize_unicode_text(item.block_type),
            text=sanitize_unicode_text(item.text),
            order=item.order,
            section_key=sanitize_unicode_text(item.section_key),
            section_title=sanitize_unicode_text(item.section_title),
            level=item.level,
            heading_path=tuple(
                sanitize_unicode_text(value) for value in item.heading_path
            ),
            page=item.page,
            metadata=sanitize_unicode_data(item.metadata),
        )
        for item in parsed.blocks
    ]
    quality = sanitize_unicode_data(parsed.quality)
    warnings = list(quality.get("warnings") or [])
    warning = f"解析文本中 {unsafe_count} 个非法 Unicode 码元已安全替换"
    if warning not in warnings:
        warnings.append(warning)
    quality.update({
        "warnings": warnings,
        "unicode_sanitized_code_units": int(
            quality.get("unicode_sanitized_code_units") or 0
        ) + unsafe_count,
    })
    return ParsedDocument(
        blocks=blocks,
        parser_type=sanitize_unicode_text(parsed.parser_type),
        parser_version=sanitize_unicode_text(parsed.parser_version),
        quality=quality,
    )


def docling_available() -> bool:
    return Config.docling_enabled and importlib.util.find_spec("docling") is not None


def validate_processing_request(filename: str, parser: str = "auto", strategy: str = "auto") -> None:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in Config.allowed_suffixes:
        supported = "、".join(sorted(Config.allowed_suffixes))
        raise ValueError(f"不支持 {suffix or '无后缀'} 文件；当前支持：{supported}")
    if parser not in PARSER_OPTIONS:
        raise ValueError(f"未知解析器：{parser}")
    if strategy not in CHUNK_STRATEGY_OPTIONS:
        raise ValueError(f"未知分块策略：{strategy}")
    if parser != "auto" and suffix not in _PARSER_SUFFIXES[parser]:
        raise ValueError(f"解析器“{PARSER_OPTIONS[parser]}”不能处理 {suffix} 文件")
    if strategy != "auto" and suffix not in _STRATEGY_SUFFIXES[strategy]:
        raise ValueError(f"分块策略“{CHUNK_STRATEGY_OPTIONS[strategy]}”不适用于 {suffix} 文件")
    if parser == "docling" and not docling_available():
        raise RuntimeError("已选择 Docling，但当前环境未安装可选依赖；请安装 requirements-docling.txt")


def _heading_info(paragraph: Paragraph) -> tuple[bool, int, str]:
    text = paragraph.text.strip()
    style_name = getattr(paragraph.style, "name", "") or ""
    match = re.search(r"(?:Heading|标题)\s*([1-9])", style_name, flags=re.IGNORECASE)
    if match and text:
        return True, int(match.group(1)), text
    numbered = _numbered_heading(text)
    if numbered:
        return True, min(numbered[0].count(".") + 1, 6), text
    return False, 0, text


def _numbered_heading(text: str) -> tuple[str, str] | None:
    """Return a conservative numbered-heading match.

    Plain PDF text commonly starts each table row with an integer. Treating
    every such row as a heading creates one relation-extraction chunk per row
    and loses the row's shared context. Real headings are short and do not use
    multi-column spacing or table-of-contents dot leaders.
    """

    match = (
        _DOTTED_NUMBERED_HEADING.match(text)
        or _SPACED_DOTTED_NUMBERED_HEADING.match(text)
        or _SINGLE_NUMBERED_HEADING.match(text)
    )
    if match is None or len(text) > 140:
        return None
    number = re.sub(r"\s+", "", match.group(1))
    title = match.group(2).strip()
    title_tokens = title.split()
    first_token = title_tokens[0] if title_tokens else ""
    single_integer = "." not in number
    integer_value = int(number) if single_integer else 0
    compact_column_flags = sum(
        token.upper() in {"Y", "N", "YES", "NO", "TRUE", "FALSE", "-"}
        for token in title_tokens
    )
    call_like_identifier = bool(re.fullmatch(
        r"[A-Za-z_][A-Za-z0-9_.:-]*\([^)]*\)", first_token,
    ))
    unmatched_delimiter = sum(title.count(value) for value in "([（【") > sum(
        title.count(value) for value in ")]）】"
    )
    looks_like_numbered_table_row = bool(
        single_integer
        and (
            any("_" in token for token in title_tokens)
            or compact_column_flags >= 2
            or (call_like_identifier and len(title_tokens) >= 2)
            or (
                integer_value >= 4
                and len(title_tokens) == 1
                and first_token.isascii()
                and first_token.isalpha()
                and first_token.upper() == first_token
            )
            or (integer_value >= 20 and unmatched_delimiter)
            or integer_value >= 100
            or (
                len(title_tokens) >= 2
                and any(
                    len(token) >= 2
                    and token.isascii()
                    and any(value.isalpha() for value in token)
                    and token.upper() == token
                    for token in title_tokens
                )
                and (
                    len(title_tokens) >= 3
                    or bool(re.search(
                        r"[\u3400-\u4dbf\u4e00-\u9fff]",
                        " ".join(title_tokens[1:]),
                    ))
                )
            )
        )
    )
    if (
        not title
        # PDF table decimals are often flattened as ``87 . 10``. A heading
        # requires a lexical title; otherwise ``87`` becomes a synthetic
        # section and unrelated cells acquire a false proposition boundary.
        or not re.search(r"[A-Za-z\u3400-\u4dbf\u4e00-\u9fff]", title)
        or len(title) > 96
        or _TOC_DOT_LEADER.search(title)
        or re.search(r"\S\s{2,}\S", title)
        or len(title_tokens) > 8
        or looks_like_numbered_table_row
    ):
        return None
    return number, title


def _docling_heading_is_likely_body_fragment(title: str) -> bool:
    """Reject high-confidence layout headings without deleting their text.

    Docling can occasionally classify a visually prominent table value or a
    wrapped sentence fragment as a Markdown heading.  Keep this deliberately
    narrow and apply it only to Docling-generated Markdown: native Markdown
    headings are author-owned structure and must remain untouched.
    """

    normalized = re.sub(r"[ \t\u3000]+", " ", title.strip())
    if not normalized:
        return True
    # A prose fragment ending in a full stop is not a section title.  Colons,
    # question marks and exclamation marks remain valid for authored headings.
    if len(normalized) <= 96 and normalized.endswith(("。", ".")):
        return True
    # Two-or-more decimal places plus a measurement unit is a strong signal
    # for a table/body value (for example latency or memory consumption), not
    # a document hierarchy marker.  Preserve it as ordinary paragraph text.
    return bool(_DOCLING_DECIMAL_MEASUREMENT_HEADING.fullmatch(normalized))


def _slug(value: str, index: int) -> str:
    safe = re.sub(r"[^\w\u4e00-\u9fff]+", "_", value)[:48].strip("_") or "section"
    return f"s{index:04d}_{safe}"


def _paper_score(text: str) -> int:
    return len({item.group(0).lower() for item in _PAPER_MARKERS.finditer(text[:60000])})


def inspect_document(path: Path) -> DocumentProfile:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF 解析依赖未安装，请执行 pip install -r requirements.txt") from exc
        reader = PdfReader(str(path))
        sample_count = min(len(reader.pages), max(1, Config.pdf_probe_pages))
        texts = [(reader.pages[index].extract_text() or "") for index in range(sample_count)]
        sample = "\n".join(texts)
        empty = sum(not text.strip() for text in texts)
        replacement = sample.count("�")
        table_lines = sum(1 for line in sample.splitlines() if line.count("|") >= 2 or re.search(r"\S\s{3,}\S", line))
        complexity = "complex" if table_lines >= 4 or replacement / max(1, len(sample)) > 0.01 else "simple"
        return DocumentProfile(
            suffix=suffix,
            page_count=len(reader.pages),
            sampled_pages=sample_count,
            text_chars=len(sample.strip()),
            has_text_layer=bool(sample.strip()),
            empty_page_ratio=round(empty / max(1, sample_count), 4),
            replacement_ratio=round(replacement / max(1, len(sample)), 6),
            table_count=1 if table_lines >= 4 else 0,
            paper_score=_paper_score(sample),
            layout_complexity=complexity,
            metadata={"sample_table_lines": table_lines},
        )
    if suffix == ".docx":
        document = Document(path)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        headings = sum(_heading_info(p)[0] for p in document.paragraphs)
        images = len(document.inline_shapes)
        text = "\n".join(paragraphs)
        complexity = "complex" if document.tables or images > 4 else "simple"
        return DocumentProfile(suffix=suffix, text_chars=len(text), heading_count=headings,
                               table_count=len(document.tables), image_count=images,
                               paper_score=_paper_score(text), layout_complexity=complexity)
    if suffix in {".xlsx", ".xls", ".csv"}:
        return DocumentProfile(suffix=suffix, table_count=1, layout_complexity="tabular")
    text = path.read_text(encoding="utf-8", errors="ignore")
    headings = len(re.findall(r"^#{1,6}\s+", text, flags=re.MULTILINE))
    return DocumentProfile(suffix=suffix, text_chars=len(text.strip()), heading_count=headings,
                           paper_score=_paper_score(text))


def route_processing(path: Path, requested_parser: str = "auto", requested_strategy: str = "auto") -> ProcessingPlan:
    validate_processing_request(path.name, requested_parser, requested_strategy)
    profile = inspect_document(path)
    suffix = profile.suffix
    reasons: list[str] = []
    parser = requested_parser
    if parser == "auto":
        if suffix == ".docx":
            parser = "native_docx"; reasons.append("DOCX 使用原生结构解析")
        elif suffix in {".md", ".markdown", ".txt"}:
            parser = "native_text"; reasons.append("文本文件使用标题与段落解析")
        elif suffix in {".xlsx", ".xls", ".csv"}:
            parser = "spreadsheet"; reasons.append("电子表格使用工作表与行解析")
        elif not profile.has_text_layer:
            if not docling_available():
                raise RuntimeError("PDF 未检测到可用文本层，且当前环境未安装 Docling，无法可靠解析扫描件")
            parser = "docling"; reasons.append("PDF 无文本层，改用 Docling")
        elif profile.layout_complexity == "complex" and docling_available():
            parser = "docling"; reasons.append("PDF 版面较复杂，改用 Docling")
        else:
            parser = "plaintext_pdf"; reasons.append("PDF 文本层完整，使用轻量解析")
    else:
        reasons.append(f"按用户选择使用{PARSER_OPTIONS[parser]}")

    strategy = requested_strategy
    if strategy == "auto":
        if suffix in {".xlsx", ".xls", ".csv"}:
            strategy = "table"; reasons.append("表格文件按行构造知识块")
        elif profile.paper_score >= 4:
            strategy = "paper"; reasons.append("检测到摘要/章节/参考文献等论文结构")
        else:
            strategy = "general"; reasons.append("采用通用结构与 token 预算分块")
    else:
        reasons.append(f"按用户选择使用{CHUNK_STRATEGY_OPTIONS[strategy]}")
    return ProcessingPlan(requested_parser, requested_strategy, parser, strategy, "；".join(reasons), profile)


def _heading_path(stack: list[str], level: int, title: str) -> tuple[str, ...]:
    del stack[max(0, level - 1):]
    while len(stack) < max(0, level - 1):
        stack.append("")
    stack.append(title)
    return tuple(item for item in stack if item)


def _table_text(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    return "\n".join([
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
        *("| " + " | ".join(row) + " |" for row in normalized[1:]),
    ])


def _table_header_projection(rows: list[list[str]]) -> tuple[int, list[str]]:
    """Return multi-row header depth and readable per-column labels."""

    if not rows:
        return 0, []
    width = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in rows]

    def normalized_cell(value: Any) -> str:
        return re.sub(
            r"[^0-9a-z\u3400-\u4dbf\u4e00-\u9fff]+", "",
            unicodedata.normalize("NFKC", str(value or "")).casefold(),
        )

    def numeric_measurement(value: Any) -> bool:
        text = unicodedata.normalize(
            "NFKC", str(value or ""),
        ).strip().replace(" ", "").strip("*`_")
        return bool(re.fullmatch(
            r"[<>≤≥~≈]?[+-]?\d+(?:\.\d+)?"
            r"(?:[±+-]\d+(?:\.\d+)?)?%?(?:\([^)]*\))?",
            text,
        ))

    primary_key = normalized_cell(normalized_rows[0][0])
    depth = 1
    for candidate in normalized_rows[1:4]:
        candidate_key = normalized_cell(candidate[0])
        sublabels = [value.strip() for value in candidate[1:] if value.strip()]
        if (
            candidate_key not in {"", primary_key}
            or not sublabels
            or any(numeric_measurement(value) for value in sublabels)
        ):
            break
        depth += 1

    headers: list[str] = []
    for column_index in range(width):
        parts: list[str] = []
        seen: set[str] = set()
        for row in normalized_rows[:depth]:
            value = str(row[column_index] or "").strip()
            identity = normalized_cell(value)
            if not value or not identity or identity in seen:
                continue
            seen.add(identity)
            parts.append(value)
        headers.append(" / ".join(parts))
    return depth, headers


def _append_block(blocks: list[DocumentBlock], block_type: str, text: str, section_key: str,
                  section_title: str, level: int, path: tuple[str, ...], *, page: int | None = None,
                  metadata: dict[str, Any] | None = None) -> None:
    text = text.strip()
    if not text:
        return
    order = len(blocks) + 1
    blocks.append(DocumentBlock(f"b{order:05d}", block_type, text, order, section_key, section_title,
                                level, path, page, metadata or {}))


def _parse_docx(path: Path) -> ParsedDocument:
    document = Document(path)
    blocks: list[DocumentBlock] = []
    title, level, section_index = path.stem, 1, 1
    section_key = _slug(title, section_index)
    stack = [title]
    table_index = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            is_heading, candidate_level, text = _heading_info(item)
            if not text:
                continue
            if is_heading:
                title, level = text, candidate_level
                section_index += 1
                section_key = _slug(title, section_index)
                path_parts = _heading_path(stack, level, title)
                _append_block(blocks, "heading", title, section_key, title, level, path_parts)
            else:
                _append_block(blocks, "paragraph", text, section_key, title, level, tuple(stack))
        elif isinstance(item, Table):
            table_index += 1
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in item.rows]
            table = _table_text(rows)
            _append_block(blocks, "table", table, section_key, title, level, tuple(stack),
                          metadata={"table_index": table_index, "rows": len(rows)})
            header_depth, headers = _table_header_projection(rows)
            for row_index, values in enumerate(rows[header_depth:], start=1):
                pairs = [f"{headers[i] or f'列{i + 1}'}: {value}" for i, value in enumerate(values) if value]
                _append_block(blocks, "table_row", "；".join(pairs), section_key, title, level, tuple(stack),
                              metadata={"table_index": table_index, "row_index": row_index,
                                        "headers": headers, "values": values})
    return _parsed(
        blocks,
        "native_docx",
        f"python-docx-{DOCUMENT_STRUCTURE_PROFILE_VERSION}",
        allow_empty=True,
    )


def _text_to_blocks(
    text: str,
    stem: str,
    *,
    page: int | None = None,
    initial_order: int = 0,
    page_break_marker: str = "",
    validate_layout_headings: bool = False,
) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    title, level, section_index = stem, 1, 1
    section_key = _slug(title, section_index)
    stack = [title]
    paragraph: list[str] = []
    table_lines: list[str] = []
    current_page = page

    def flush_paragraph() -> None:
        if paragraph:
            _append_block(blocks, "paragraph", "\n".join(paragraph), section_key, title, level,
                          tuple(stack), page=current_page)
            paragraph.clear()

    def flush_table() -> None:
        if table_lines:
            _append_block(blocks, "table", "\n".join(table_lines), section_key, title, level,
                          tuple(stack), page=current_page)
            rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
            rows = [row for row in rows if not all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in row)]
            header_depth, headers = _table_header_projection(rows)
            for row_index, values in enumerate(rows[header_depth:], start=1):
                pairs = [f"{headers[i] or f'列{i + 1}'}: {value}" for i, value in enumerate(values) if value]
                _append_block(blocks, "table_row", "；".join(pairs), section_key, title, level,
                              tuple(stack), page=current_page,
                              metadata={"row_index": row_index, "headers": headers, "values": values})
            table_lines.clear()

    for raw_line in text.splitlines():
        # Leading whitespace is source structure for nested lists.  Earlier
        # revisions stripped every line before assembling a paragraph, which
        # flattened children and siblings into one level before the relation
        # pipeline could inspect them.  Use a stripped view only for routing;
        # retain the original indentation in paragraph text.
        line = raw_line.rstrip()
        logical_line = line.strip()
        if page_break_marker and logical_line == page_break_marker:
            flush_paragraph()
            flush_table()
            current_page = int(current_page or 0) + 1
            continue
        markdown = re.match(r"^(#{1,6})\s+(.+)$", logical_line)
        demoted_layout_heading = bool(
            markdown
            and validate_layout_headings
            and _docling_heading_is_likely_body_fragment(markdown.group(2))
        )
        if demoted_layout_heading:
            # Strip the exporter-added Markdown marker but retain the source
            # text in the surrounding paragraph and on the original page.
            line = markdown.group(2).strip()
            logical_line = line
            markdown = None
        numbered = None if demoted_layout_heading else _numbered_heading(logical_line)
        if (
            numbered
            and validate_layout_headings
            and _docling_heading_is_likely_body_fragment(logical_line)
        ):
            # Docling may emit the same visual heading either as Markdown or
            # as an unmarked line that resembles a numbered title.  Apply the
            # same layout-only sanity rule before the numbered-heading route.
            numbered = None
        is_table = (
            logical_line.startswith("|")
            and logical_line.endswith("|")
            and logical_line.count("|") >= 2
        )
        if markdown or numbered:
            flush_paragraph(); flush_table()
            if markdown:
                level, title = len(markdown.group(1)), markdown.group(2).strip()
            else:
                level, title = (
                    min(numbered[0].count(".") + 1, 6), logical_line,
                )
            section_index += 1
            section_key = _slug(title, section_index)
            path_parts = _heading_path(stack, level, title)
            _append_block(
                blocks, "heading", title, section_key, title, level,
                path_parts, page=current_page,
            )
        elif is_table:
            flush_paragraph(); table_lines.append(logical_line)
        elif not logical_line:
            flush_paragraph(); flush_table()
        else:
            flush_table(); paragraph.append(line)
    flush_paragraph(); flush_table()
    if initial_order:
        blocks = [DocumentBlock(f"b{initial_order + i:05d}", item.block_type, item.text,
                                initial_order + i, item.section_key, item.section_title, item.level,
                                item.heading_path, item.page, item.metadata)
                  for i, item in enumerate(blocks, start=1)]
    return blocks


def _parse_text(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return _parsed(
        _text_to_blocks(text, path.stem),
        "native_text",
        f"builtin-text-{DOCUMENT_STRUCTURE_PROFILE_VERSION}",
    )


def _parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    blocks: list[DocumentBlock] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        page_blocks = _text_to_blocks(text, f"第 {index} 页", page=index, initial_order=len(blocks))
        blocks.extend([
            DocumentBlock(item.key, item.block_type, item.text, item.order,
                          f"p{index:04d}_{item.section_key}", item.section_title, item.level,
                          item.heading_path, item.page, item.metadata)
            for item in page_blocks
        ])
    return _parsed(blocks, "plaintext_pdf", (
        f"pypdf-{_package_version('pypdf')}-{DOCUMENT_STRUCTURE_PROFILE_VERSION}"
    ),
                   {"page_count": len(reader.pages)}, allow_empty=True)


def _docling_cuda_memory() -> tuple[bool, int, int, str]:
    """Return CUDA readiness and free/total MiB without making CUDA mandatory."""
    try:
        import torch
        if not torch.cuda.is_available():
            return False, 0, 0, "容器未检测到可用 CUDA"
        free_bytes, total_bytes = torch.cuda.mem_get_info(0)
        free_mb = int(free_bytes / 1024 / 1024)
        total_mb = int(total_bytes / 1024 / 1024)
        if free_mb < Config.docling_gpu_min_free_mb:
            return False, free_mb, total_mb, (
                f"剩余显存 {free_mb} MiB 低于安全门槛 {Config.docling_gpu_min_free_mb} MiB"
            )
        return True, free_mb, total_mb, ""
    except Exception as exc:
        return False, 0, 0, f"CUDA 状态检查失败：{exc}"


def _select_docling_device() -> tuple[str, dict[str, Any]]:
    requested = Config.docling_device or "auto"
    metadata: dict[str, Any] = {"docling_accelerator_requested": requested}
    if requested == "cpu":
        return "cpu", metadata
    if requested not in {"auto", "cuda"} and not requested.startswith("cuda:"):
        raise ValueError("DOCLING_DEVICE 仅支持 auto、cpu、cuda 或 cuda:N")

    ready, free_mb, total_mb, reason = _docling_cuda_memory()
    metadata.update({"docling_gpu_free_mb_before": free_mb, "docling_gpu_total_mb": total_mb})
    if ready:
        return "cuda:0" if requested == "auto" else requested, metadata
    if not Config.docling_cuda_fallback and requested.startswith("cuda"):
        raise RuntimeError(reason)
    metadata["docling_gpu_fallback_reason"] = reason
    return "cpu", metadata


def _is_cuda_failure(exc: BaseException) -> bool:
    message = f"{type(exc).__name__}: {exc}".lower()
    markers = (
        "cuda", "cudnn", "cublas", "out of memory", "device-side assert",
        "no kernel image", "nvidia driver", "hip out of memory",
    )
    return any(marker in message for marker in markers)


def _release_docling_cuda() -> None:
    gc.collect()
    try:
        import torch
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
            torch.cuda.ipc_collect()
    except Exception:
        pass


_DOCLING_PAGE_BREAK_MARKER = "<!-- knowledge-forge-page-break -->"


def _docling_markdown(path: Path, device: str) -> str:
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import AcceleratorOptions, PdfPipelineOptions, RapidOcrOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    batch_size = Config.docling_batch_size
    accelerator = AcceleratorOptions(
        num_threads=Config.docling_num_threads,
        device=device,
        cuda_use_flash_attention2=False,
    )
    pipeline = PdfPipelineOptions(
        accelerator_options=accelerator,
        ocr_options=RapidOcrOptions(backend="torch"),
        ocr_batch_size=batch_size,
        layout_batch_size=batch_size,
        table_batch_size=batch_size,
        queue_max_size=max(4, batch_size * 4),
    )
    converter = DocumentConverter(
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline)}
    )
    try:
        result = converter.convert(str(path))
        return result.document.export_to_markdown(
            page_break_placeholder=_DOCLING_PAGE_BREAK_MARKER,
        )
    finally:
        del converter


def _parse_docling(path: Path) -> ParsedDocument:
    if not docling_available():
        raise RuntimeError("Docling 未安装；请安装 requirements-docling.txt 后重试")
    with _DOCLING_LOCK:
        device, accelerator_metadata = _select_docling_device()
        if device.startswith("cuda"):
            try:
                import torch
                fraction = min(1.0, max(0.05, Config.docling_gpu_memory_fraction))
                torch.cuda.set_per_process_memory_fraction(fraction, 0)
                accelerator_metadata["docling_gpu_memory_fraction"] = fraction
            except Exception as exc:
                if not Config.docling_cuda_fallback:
                    raise
                accelerator_metadata["docling_gpu_fallback_reason"] = f"显存限额配置失败：{exc}"
                device = "cpu"
        LOGGER.info("Docling parsing %s with %s", path.name, device)
        try:
            markdown = _docling_markdown(path, device)
        except Exception as exc:
            if not (device.startswith("cuda") and Config.docling_cuda_fallback and _is_cuda_failure(exc)):
                raise
            LOGGER.warning("Docling CUDA failed for %s; retrying on CPU: %s", path.name, exc)
            accelerator_metadata["docling_gpu_fallback_reason"] = f"CUDA 解析失败：{exc}"
            _release_docling_cuda()
            device = "cpu"
            markdown = _docling_markdown(path, device)
        finally:
            if device.startswith("cuda"):
                _release_docling_cuda()
    accelerator_metadata.update({
        "docling_accelerator": device,
        "docling_batch_size": Config.docling_batch_size,
    })
    return _parsed(_text_to_blocks(
        markdown,
        path.stem,
        page=1,
        page_break_marker=_DOCLING_PAGE_BREAK_MARKER,
        validate_layout_headings=True,
    ), "docling",
                   f"docling-{_package_version('docling')}-{DOCUMENT_STRUCTURE_PROFILE_VERSION}",
                   accelerator_metadata,
                   allow_empty=path.suffix.lower() in {".pdf", ".docx"})


def _sheet_blocks(sheet_name: str, rows: Iterable[list[Any]], blocks: list[DocumentBlock]) -> None:
    cleaned = [["" if value is None else str(value).strip() for value in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return
    section_key = _slug(sheet_name, len({item.section_key for item in blocks}) + 1)
    header_depth, headers = _table_header_projection(cleaned)
    _append_block(blocks, "heading", sheet_name, section_key, sheet_name, 1, (sheet_name,),
                  metadata={"sheet": sheet_name})
    _append_block(blocks, "table", _table_text(cleaned), section_key, sheet_name, 1, (sheet_name,),
                  metadata={"sheet": sheet_name, "rows": len(cleaned)})
    for row_index, values in enumerate(cleaned[header_depth:], start=1):
        pairs = [f"{headers[i] or f'列{i + 1}'}: {value}" for i, value in enumerate(values) if value]
        _append_block(blocks, "table_row", "；".join(pairs), section_key, sheet_name, 1, (sheet_name,),
                      metadata={"sheet": sheet_name, "row_index": row_index,
                                "headers": headers, "values": values})


def _parse_spreadsheet(path: Path) -> ParsedDocument:
    blocks: list[DocumentBlock] = []
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig", errors="ignore")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|") if text.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        _sheet_blocks(path.stem, list(csv.reader(text.splitlines(), dialect)), blocks)
        version = "csv-stdlib"
    elif suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("XLSX 解析依赖未安装，请执行 pip install -r requirements.txt") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            _sheet_blocks(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)], blocks)
        version = f"openpyxl-{_package_version('openpyxl')}"
    else:
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("XLS 解析依赖未安装，请执行 pip install -r requirements.txt") from exc
        workbook = xlrd.open_workbook(path)
        for sheet in workbook.sheets():
            _sheet_blocks(sheet.name, [sheet.row_values(index) for index in range(sheet.nrows)], blocks)
        version = f"xlrd-{_package_version('xlrd')}"
    return _parsed(blocks, "spreadsheet", version)


def _package_version(name: str) -> str:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return "unknown"


def _parsed(blocks: list[DocumentBlock], parser: str, version: str,
            extra: dict[str, Any] | None = None, *, allow_empty: bool = False) -> ParsedDocument:
    text_chars = sum(len(item.text) for item in blocks if item.block_type != "heading")
    quality = {
        "text_chars": text_chars,
        "block_count": len(blocks),
        "heading_count": sum(item.block_type == "heading" for item in blocks),
        "table_count": sum(item.block_type == "table" for item in blocks),
        "table_row_count": sum(item.block_type == "table_row" for item in blocks),
        "warnings": [],
        **(extra or {}),
    }
    if (not blocks or text_chars == 0) and not allow_empty:
        raise ValueError("解析完成但未提取到可入库文本")
    return ParsedDocument(blocks, parser, version, quality)


def parse_document(path: Path, parser: str) -> ParsedDocument:
    return {
        "native_docx": _parse_docx,
        "native_text": _parse_text,
        "plaintext_pdf": _parse_pdf,
        "docling": _parse_docling,
        "spreadsheet": _parse_spreadsheet,
    }[parser](path)


def blocks_to_sections(blocks: list[DocumentBlock]) -> list[Section]:
    grouped: dict[str, dict[str, Any]] = {}
    for block in blocks:
        if block.block_type in {"heading", "table_row"}:
            continue
        item = grouped.setdefault(block.section_key, {
            "title": block.section_title,
            "level": block.level,
            "parts": [],
            "pages": set(),
            "block_types": set(),
            "heading_path": list(block.heading_path),
            "media_assets": [],
        })
        item["parts"].append({
            "text": block.text,
            "graph_extraction_enabled": (
                block.metadata.get("graph_extraction_enabled") is not False
            ),
            "graph_exclusion_reason": (
                "non_substantive_document_image"
                if block.block_type == "image"
                and block.metadata.get("graph_extraction_enabled") is False
                else ""
            ),
        })
        item["block_types"].add(block.block_type)
        if block.page is not None:
            item["pages"].add(block.page)
        if block.block_type == "image" and block.metadata:
            item["media_assets"].append(dict(block.metadata))
    sections = []
    for key, item in grouped.items():
        content = ""
        graph_exclusion_spans: list[dict[str, Any]] = []
        for part in item["parts"]:
            text = str(part.get("text") or "").strip()
            if not text:
                continue
            if content:
                content += "\n\n"
            start = len(content)
            content += text
            end = len(content)
            if part.get("graph_extraction_enabled") is False:
                graph_exclusion_spans.append({
                    "start": start,
                    "end": end,
                    "reason": str(
                        part.get("graph_exclusion_reason")
                        or "retrieval_only_source_region"
                    ),
                })
        if content:
            sections.append(Section(key, item["title"], item["level"], content, {
                "pages": sorted(item["pages"]),
                "block_types": sorted(item["block_types"]),
                "heading_path": item["heading_path"],
                "media_assets": item["media_assets"],
                "graph_exclusion_spans": graph_exclusion_spans,
                "graph_extraction_enabled": any(
                    part.get("graph_extraction_enabled") is not False
                    for part in item["parts"]
                ),
            }))
    return sections


def extract_docx_sections(path: Path) -> list[Section]:
    return blocks_to_sections(_parse_docx(path).blocks)


def extract_plain_sections(path: Path) -> list[Section]:
    parsed = _parse_pdf(path) if path.suffix.lower() == ".pdf" else _parse_text(path)
    return blocks_to_sections(parsed.blocks)


def extract_sections(path: Path) -> list[Section]:
    if path.suffix.lower() == ".docx":
        return extract_docx_sections(path)
    return extract_plain_sections(path)


def estimate_tokens(text: str) -> int:
    cjk = len(re.findall(r"[\u3400-\u9fff]", text))
    words = len(re.findall(r"[A-Za-z0-9_]+", text))
    other = len(re.findall(r"[^\s\u3400-\u9fffA-Za-z0-9_]", text))
    return max(1, cjk + math.ceil(words * 1.3) + math.ceil(other / 4))


def _segment_spans(text: str) -> list[tuple[int, int, str]]:
    """Return shared sentence spans without normalizing away source coordinates."""

    return split_source_sentences(text)


def _segments(text: str) -> list[str]:
    # Sentence splitting treats a newline as a boundary.  Leading whitespace
    # on the following span is therefore source structure, not cosmetic
    # padding: it carries nested-list depth into the final chunk.  Stripping
    # every span flattened parent/child lists even though the parser blocks
    # still retained their indentation.  Remove only trailing separators and
    # keep the original left margin.
    return [value.rstrip() for _, _, value in _segment_spans(text) if value.strip()]


def _token_pieces(text: str, target_tokens: int, max_tokens: int, overlap_percent: int) -> list[str]:
    segments: list[str] = []
    for segment in _segments(text):
        tokens = estimate_tokens(segment)
        if tokens <= max_tokens:
            segments.append(segment)
            continue
        chars_per_token = max(1.0, len(segment) / tokens)
        width = max(1, int(max_tokens * chars_per_token * 0.95))
        segments.extend(segment[start:start + width] for start in range(0, len(segment), width))
    pieces: list[str] = []
    current: list[str] = []
    current_tokens = 0
    for segment in segments:
        segment_tokens = estimate_tokens(segment)
        if current and current_tokens + segment_tokens > max_tokens:
            pieces.append("\n".join(current).strip())
            overlap_budget = max(0, int(target_tokens * overlap_percent / 100))
            overlap: list[str] = []
            overlap_tokens = 0
            for previous in reversed(current):
                amount = estimate_tokens(previous)
                if overlap and overlap_tokens + amount > overlap_budget:
                    break
                if amount <= overlap_budget:
                    overlap.insert(0, previous); overlap_tokens += amount
            current, current_tokens = overlap, overlap_tokens
        current.append(segment); current_tokens += segment_tokens
        if current_tokens >= target_tokens:
            pieces.append("\n".join(current).strip())
            current, current_tokens = [], 0
    if current:
        pieces.append("\n".join(current).strip())
    return [piece for piece in pieces if piece] or [""]


def _enforce_chunk_token_limit(chunks: list[TextChunk], max_tokens: int) -> list[TextChunk]:
    """Split any oversized strategy output while retaining its project metadata."""
    if max_tokens < 64:
        raise ValueError("单切片 Token 上限不能小于 64")
    bounded: list[TextChunk] = []
    parent_key_map: dict[str, str] = {}
    marker = re.compile(
        r"(?s)^(# .*?\n\n章节ID: .*?\n来源文档: .*?\n\n)分块:\s*\d+/\d+\n\n(.*)$"
    )
    for chunk in chunks:
        if chunk.token_count <= max_tokens:
            bounded.append(chunk)
            parent_key_map[chunk.key] = chunk.key
            continue

        match = marker.match(chunk.content)
        prefix, body = (match.group(1), match.group(2)) if match else ("", chunk.content)
        framing_tokens = estimate_tokens(prefix + "分块: 1/1\n\n") if prefix else 0
        payload_limit = max_tokens - framing_tokens
        if payload_limit >= 16:
            pieces = _token_pieces(body, payload_limit, payload_limit, 0)
            contents = [f"{prefix}分块: {index}/{len(pieces)}\n\n{piece}".strip()
                        for index, piece in enumerate(pieces, start=1)]
        else:
            contents = _token_pieces(chunk.content, max_tokens, max_tokens, 0)

        created: list[TextChunk] = []
        for index, content in enumerate(contents, start=1):
            # Estimation and splitting use the same tokenizer proxy, so this is a real hard bound.
            if estimate_tokens(content) > max_tokens:
                raise ValueError(f"知识块固定头信息已超过项目分块上限 {max_tokens}")
            key = f"{chunk.key}_s{index:03d}"
            created.append(TextChunk(
                key=key,
                section_key=chunk.section_key,
                title=chunk.title,
                part_index=index,
                part_total=len(contents),
                content=content,
                content_hash=hashlib.sha256(content.encode("utf-8")).hexdigest(),
                token_count=estimate_tokens(content),
                parent_key=chunk.parent_key,
                metadata={**chunk.metadata, "token_limit_split": True},
            ))
        bounded.extend(created)
        parent_key_map[chunk.key] = created[0].key

    return [
        TextChunk(
            key=item.key,
            section_key=item.section_key,
            title=item.title,
            part_index=item.part_index,
            part_total=item.part_total,
            content=item.content,
            content_hash=item.content_hash,
            token_count=item.token_count,
            parent_key=parent_key_map.get(item.parent_key, item.parent_key),
            metadata=item.metadata,
        )
        for item in bounded
    ]


def _chunk(section: Section, source_name: str, body: str, index: int, total: int, strategy: str,
           *, key_suffix: str | None = None, parent_key: str | None = None,
           metadata: dict[str, Any] | None = None) -> TextChunk:
    prefix = f"# {section.title}\n\n章节ID: {section.key}\n来源文档: {source_name}\n\n"
    content = f"{prefix}分块: {index}/{total}\n\n{body}".strip()
    key = f"{section.key}_{key_suffix or f'p{index:03d}'}"
    data = {"strategy": strategy, "heading_path": section.metadata.get("heading_path", []),
            "pages": section.metadata.get("pages", []),
            "block_types": section.metadata.get("block_types", []),
            "media_assets": section.metadata.get("media_assets", []), **(metadata or {})}
    return TextChunk(key, section.key, section.title, index, total, content,
                     hashlib.sha256(content.encode("utf-8")).hexdigest(), estimate_tokens(content),
                     parent_key, data)


def _section_graph_regions(
    section: Section,
) -> list[tuple[str, bool, dict[str, Any]]]:
    """Partition one section without losing retrieval-only source text."""

    raw_spans = section.metadata.get("graph_exclusion_spans", [])
    spans: list[tuple[int, int, dict[str, Any]]] = []
    for value in raw_spans if isinstance(raw_spans, list) else []:
        if not isinstance(value, dict):
            continue
        start = max(0, min(len(section.content), int(value.get("start", 0))))
        end = max(start, min(len(section.content), int(value.get("end", start))))
        if end > start:
            spans.append((start, end, value))
    if not spans:
        return [(section.content, True, {})]

    regions: list[tuple[str, bool, dict[str, Any]]] = []
    cursor = 0
    for start, end, value in sorted(spans, key=lambda item: item[0]):
        if start > cursor and section.content[cursor:start].strip():
            regions.append((section.content[cursor:start], True, {}))
        if end > cursor and section.content[max(cursor, start):end].strip():
            regions.append((
                section.content[max(cursor, start):end],
                False,
                {
                    "graph_exclusion_reason": str(
                        value.get("reason") or "retrieval_only_source_region"
                    ),
                    "summary_kind": str(value.get("summary_kind") or ""),
                    "summary_language": str(value.get("summary_language") or ""),
                },
            ))
        cursor = max(cursor, end)
    if cursor < len(section.content) and section.content[cursor:].strip():
        regions.append((section.content[cursor:], True, {}))
    return regions or [(section.content, True, {})]


def split_sections(sections: list[Section], *, max_chars: int, overlap_chars: int,
                   source_name: str) -> list[TextChunk]:
    """Compatibility strategy: preserve the original character-budget behavior."""
    chunks: list[TextChunk] = []
    for section in sections:
        prefix_size = len(f"# {section.title}\n\n章节ID: {section.key}\n来源文档: {source_name}\n\n")
        payload_limit = max(240, max_chars - prefix_size)
        pieces: list[tuple[str, bool, dict[str, Any]]] = []
        for region, region_graph_enabled, region_metadata in _section_graph_regions(section):
            region_pieces: list[str] = []
            current = ""
            for segment in _segments(region):
                if len(segment) > payload_limit:
                    if current:
                        region_pieces.append(current); current = ""
                    region_pieces.extend(
                        segment[start:start + payload_limit]
                        for start in range(0, len(segment), payload_limit)
                    )
                    continue
                candidate = f"{current}\n{segment}".strip()
                if current and len(candidate) > payload_limit:
                    region_pieces.append(current)
                    current = (
                        f"{current[-overlap_chars:].strip()}\n{segment}".strip()
                        if overlap_chars else segment
                    )
                else:
                    current = candidate
            if current:
                region_pieces.append(current)
            pieces.extend(
                (body, region_graph_enabled, region_metadata)
                for body in (region_pieces or [""])
            )
        section_graph_enabled = bool(
            section.metadata.get("graph_extraction_enabled", True)
        )
        chunks.extend(
            _chunk(
                section, source_name, body, index, len(pieces), "default",
                metadata={
                    "graph_extraction_enabled": (
                        section_graph_enabled and region_graph_enabled
                    ),
                    **{
                        key: section.metadata[key]
                        for key in (
                            "document_primary_language",
                            "bilingual_mirror_detection",
                        )
                        if key in section.metadata
                    },
                    **region_metadata,
                },
            )
            for index, (body, region_graph_enabled, region_metadata)
            in enumerate(pieces, start=1)
        )
    return chunks


def _general_chunks(sections: list[Section], source_name: str, *, target: int, maximum: int,
                    overlap: int, strategy: str = "general") -> list[TextChunk]:
    chunks: list[TextChunk] = []
    for section in sections:
        section_graph_enabled = bool(
            section.metadata.get("graph_extraction_enabled", True)
        )
        pieces: list[tuple[str, bool, dict[str, Any]]] = []
        for region, region_graph_enabled, region_metadata in _section_graph_regions(section):
            pieces.extend(
                (body, region_graph_enabled, region_metadata)
                for body in _token_pieces(region, target, maximum, overlap)
            )
        structure_metadata = {
            key: section.metadata[key]
            for key in (
                "structural_role", "toc_line_count", "toc_line_ratio",
                "document_primary_language", "bilingual_mirror_detection",
            )
            if key in section.metadata
        }
        parent_key = None
        if len(pieces) > 1:
            overview = _token_pieces(section.content, min(220, target), min(320, maximum), 0)[0]
            parent = _chunk(section, source_name, overview, 0, len(pieces), strategy,
                            key_suffix="parent", metadata={"chunk_role": "parent",
                            "graph_extraction_enabled": False, **structure_metadata})
            chunks.append(parent); parent_key = parent.key
        for index, (body, region_graph_enabled, region_metadata) in enumerate(
            pieces, start=1,
        ):
            chunks.append(_chunk(section, source_name, body, index, len(pieces), strategy,
                                 parent_key=parent_key, metadata={
                                     "chunk_role": "child",
                                     "graph_extraction_enabled": (
                                         section_graph_enabled
                                         and region_graph_enabled
                                     ),
                                     **region_metadata,
                                     **structure_metadata,
                                 }))
    return chunks


def _annotate_general_structure_roles(sections: list[Section]) -> list[Section]:
    """Mark high-confidence navigation sections without removing retrieval text.

    Dot-leader entries ending in page numbers are a layout signal rather than
    a domain vocabulary rule.  Requiring at least three entries and a majority
    of non-empty lines keeps ordinary prose, short cross-references and
    screenshot-style bullet lists eligible for graph extraction.  Navigation
    sections remain normal chunks for vector retrieval and provenance.
    """

    annotated: list[Section] = []
    for section in sections:
        lines = [line.strip() for line in section.content.splitlines() if line.strip()]
        toc_line_count = sum(
            bool(_TOC_ENTRY_DOT_LEADER.search(line)) for line in lines
        )
        toc_line_ratio = toc_line_count / len(lines) if lines else 0.0
        if toc_line_count < 3 or toc_line_ratio < 0.5:
            annotated.append(section)
            continue
        annotated.append(Section(
            key=section.key,
            title=section.title,
            level=section.level,
            content=section.content,
            metadata={
                **section.metadata,
                "structural_role": "table_of_contents",
                "graph_extraction_enabled": False,
                "toc_line_count": toc_line_count,
                "toc_line_ratio": round(toc_line_ratio, 4),
            },
        ))
    return annotated


def _coalesce_general_micro_sections(
    sections: list[Section], *, target_tokens: int,
) -> list[Section]:
    """Pack consecutive leaf micro-sections without dropping their headings.

    API and fault-reference PDFs can contain hundreds of tiny numbered leaves.
    One model workflow per leaf is prohibitively slow, while discarding the
    leaves harms recall. This deterministic packer retains every title,
    source-section key, page and body inside a bounded aggregate section.
    Top-level and media-bearing sections remain independent.
    """

    minimum = int(Config.general_coalesce_min_sections)
    if minimum <= 0 or len(sections) < minimum:
        return sections
    budget = max(128, int(target_tokens))
    result: list[Section] = []
    pending: list[Section] = []
    pending_tokens = 0

    def flush() -> None:
        nonlocal pending_tokens
        if not pending:
            return
        if len(pending) == 1:
            result.append(pending.pop())
            pending_tokens = 0
            return
        keys = [item.key for item in pending]
        pages = sorted({
            int(page)
            for item in pending
            for page in item.metadata.get("pages", [])
        })
        block_types = sorted({
            str(value)
            for item in pending
            for value in item.metadata.get("block_types", [])
        })
        paths = [
            [str(value) for value in item.metadata.get("heading_path", [])]
            for item in pending
        ]
        common_path: list[str] = []
        for values in zip(*paths) if paths and all(paths) else ():
            if len(set(values)) != 1:
                break
            common_path.append(values[0])
        body_parts = []
        for item in pending:
            marker_level = min(6, max(2, int(item.level) + 1))
            body_parts.append(
                f"{'#' * marker_level} {item.title}\n\n{item.content}"
            )
        digest = hashlib.sha256(
            "\n".join(keys).encode("utf-8")
        ).hexdigest()[:12]
        first = pending[0]
        result.append(Section(
            key=f"{first.key}_bundle_{digest}",
            title=f"{first.title} 等 {len(pending)} 节",
            level=min(item.level for item in pending),
            content="\n\n".join(body_parts),
            metadata={
                "pages": pages,
                "block_types": block_types,
                "heading_path": common_path,
                "media_assets": [],
                "coalesced_section_count": len(pending),
                "source_section_keys": keys,
                "source_section_titles": [item.title for item in pending],
            },
        ))
        pending.clear()
        pending_tokens = 0

    for section in sections:
        section_tokens = estimate_tokens(
            f"## {section.title}\n\n{section.content}"
        )
        media_assets = section.metadata.get("media_assets", [])
        synthetic_pdf_page = bool(
            _SYNTHETIC_PDF_PAGE_TITLE.fullmatch(section.title.strip())
        )
        independently_structured = bool(
            (int(section.level) <= 1 and not synthetic_pdf_page)
            or media_assets
            or "table" in set(section.metadata.get("block_types", []))
            or not section.metadata.get("graph_extraction_enabled", True)
            or bool(section.metadata.get("graph_exclusion_spans"))
            or section_tokens >= budget
        )
        if independently_structured:
            flush()
            result.append(section)
            continue
        if pending and pending_tokens + section_tokens > budget:
            flush()
        pending.append(section)
        pending_tokens += section_tokens
    flush()
    return result


def _script_information_units(text: str) -> tuple[int, int]:
    """Return comparable CJK-character and Latin-word information units."""

    cjk = len(re.findall(r"[\u3400-\u9fff]", text))
    latin = math.ceil(len(re.findall(r"[A-Za-z]+", text)) * 1.3)
    return cjk, latin


def _language_dominates(text: str, language: str, *, minimum: int = 2) -> bool:
    """Return whether one script clearly dominates a parser fragment."""

    cjk, latin = _script_information_units(text)
    if language == "en":
        return latin >= max(minimum, math.ceil(cjk * 1.8))
    return cjk >= max(minimum, math.ceil(latin * 1.35))


def _section_pages(section: Section) -> tuple[int, ...]:
    pages: list[int] = []
    for value in section.metadata.get("pages") or ():
        try:
            page = int(value)
        except (TypeError, ValueError):
            continue
        if page > 0:
            pages.append(page)
    return tuple(sorted(set(pages)))


def _trailing_language_span(
    content: str, language: str,
) -> tuple[int, int] | None:
    """Find a translated title/byline tail split into a prior PDF section."""

    records: list[tuple[int, int, str]] = []
    offset = 0
    for raw in content.splitlines(keepends=True):
        end = offset + len(raw)
        records.append((offset, end, raw.rstrip("\r\n")))
        offset = end
    if offset < len(content) or not records:
        records.append((offset, len(content), content[offset:]))

    start = len(content)
    found = False
    for line_start, _, line in reversed(records):
        value = line.strip()
        if not value:
            if found:
                start = line_start
            continue
        if not _language_dominates(value, language):
            break
        found = True
        start = line_start
    if not found or not content[start:].strip():
        return None
    return start, len(content)


def _paper_line_end(text: str, position: int) -> int:
    newline = text.find("\n", position)
    return len(text) if newline < 0 else newline + 1


def _paper_abstract_spans_for_section(
    section: Section, section_index: int,
) -> list[_BilingualSummarySpan]:
    content = section.content
    markers: list[tuple[re.Match[str], str, re.Pattern[str]]] = []
    markers.extend(
        (match, "zh", _ZH_KEYWORD_MARKER)
        for match in _ZH_ABSTRACT_MARKER.finditer(content)
    )
    markers.extend(
        (match, "en", _EN_KEYWORD_MARKER)
        for match in _EN_ABSTRACT_MARKER.finditer(content)
    )
    markers.sort(key=lambda item: item[0].start())

    if not markers and _ABSTRACT_TITLE.fullmatch(section.title.strip()):
        language = "zh" if "摘" in section.title else "en"
        return [_BilingualSummarySpan(
            section_index, 0, len(content), language, content.strip(),
        )]

    spans: list[_BilingualSummarySpan] = []
    for marker_index, (marker, language, keyword_pattern) in enumerate(markers):
        next_marker_start = (
            markers[marker_index + 1][0].start()
            if marker_index + 1 < len(markers)
            else len(content)
        )
        keyword = keyword_pattern.search(content, marker.end(), next_marker_start)
        abstract_end = keyword.start() if keyword else next_marker_start
        span_end = (
            _paper_line_end(content, keyword.start())
            if keyword else next_marker_start
        )
        spans.append(_BilingualSummarySpan(
            section_index,
            marker.start(),
            span_end,
            language,
            content[marker.end():abstract_end].strip(),
        ))
    return spans


def _dominant_body_language(
    sections: list[Section], summary_spans: list[_BilingualSummarySpan],
) -> str | None:
    spans_by_section: dict[int, list[_BilingualSummarySpan]] = {}
    for span in summary_spans:
        spans_by_section.setdefault(span.section_index, []).append(span)

    body_parts: list[str] = []
    for section_index, section in enumerate(sections):
        if _REFERENCE_TITLE.fullmatch(section.title.strip()):
            continue
        excluded_spans = [
            (span.start, span.end)
            for span in spans_by_section.get(section_index, [])
        ]
        excluded_spans.extend(bibliography_reference_spans(section.content))
        cursor = 0
        for start, end in sorted(excluded_spans):
            if end <= cursor:
                continue
            body_parts.append(section.content[cursor:max(cursor, start)])
            cursor = max(cursor, end)
        body_parts.append(section.content[cursor:])

    cjk, latin = _script_information_units("\n".join(body_parts))
    total = cjk + latin
    if total < 240:
        return None
    if cjk >= 200 and cjk / total >= 0.72:
        return "zh"
    if latin >= 200 and latin / total >= 0.72:
        return "en"
    return None


def _extend_secondary_abstract_frontmatter(
    content: str, span: _BilingualSummarySpan,
) -> _BilingualSummarySpan:
    """Include an adjacent translated title/byline without crossing scripts."""

    if span.start <= 0:
        return span
    line_ranges: list[tuple[int, int, str]] = []
    offset = 0
    for line in content.splitlines(keepends=True):
        line_ranges.append((offset, offset + len(line), line))
        offset += len(line)
    if offset < len(content):
        line_ranges.append((offset, len(content), content[offset:]))

    start = span.start
    blank_lines = 0
    for line_start, line_end, line in reversed(line_ranges):
        if line_end > span.start:
            continue
        if span.start - line_start > 1600:
            break
        value = line.strip()
        if not value:
            if blank_lines >= 1:
                break
            blank_lines += 1
            start = line_start
            continue
        cjk, latin = _script_information_units(value)
        secondary_dominant = (
            latin >= max(2, math.ceil(cjk * 1.8))
            if span.language == "en"
            else cjk >= max(2, math.ceil(latin * 1.35))
        )
        if not secondary_dominant:
            break
        start = line_start

    return _BilingualSummarySpan(
        span.section_index, start, span.end, span.language, span.abstract_text,
        span.kind,
    )


def _summary_heading_info(title: str) -> tuple[str, str] | None:
    normalized = re.sub(r"[ \t\u3000]+", " ", title.strip())
    for pattern, language, kind in _BILINGUAL_SUMMARY_TITLES:
        if pattern.fullmatch(normalized):
            return language, kind
    return None


def _general_summary_heading_spans(
    sections: list[Section],
) -> list[_BilingualSummarySpan]:
    spans: list[_BilingualSummarySpan] = []
    for section_index, section in enumerate(sections):
        info = _summary_heading_info(section.title)
        if info is None or not section.content.strip():
            continue
        language, kind = info
        spans.append(_BilingualSummarySpan(
            section_index, 0, len(section.content), language,
            section.content.strip(), kind,
        ))
    return spans


def _select_secondary_mirror_summaries(
    sections: list[Section], summary_spans: list[_BilingualSummarySpan], *,
    frontmatter_sections: int,
    max_section_gap: int = 4,
    frontmatter_page_limit: int | None = None,
) -> tuple[str | None, dict[int, list[_BilingualSummarySpan]]]:
    """Select conservative cross-language mirror summaries."""

    primary_language = _dominant_body_language(sections, summary_spans)
    if primary_language is None:
        return None, {}

    frontmatter_limit = min(len(sections), frontmatter_sections)

    def frontmatter_eligible(span: _BilingualSummarySpan) -> bool:
        pages = _section_pages(sections[span.section_index])
        if pages and frontmatter_page_limit is not None:
            return min(pages) <= frontmatter_page_limit
        return span.section_index < frontmatter_limit

    primary_spans = [
        span for span in summary_spans
        if span.language == primary_language
        and frontmatter_eligible(span)
        and estimate_tokens(span.abstract_text) >= 25
    ]
    secondary_spans = [
        span for span in summary_spans
        if span.language != primary_language
        and frontmatter_eligible(span)
        and estimate_tokens(span.abstract_text) >= 25
    ]
    if not primary_spans or not secondary_spans:
        return primary_language, {}

    selected: dict[int, list[_BilingualSummarySpan]] = {}
    for secondary in secondary_spans:
        compatible_primary = [
            span for span in primary_spans if span.kind == secondary.kind
        ]
        if not compatible_primary:
            continue
        primary = min(
            compatible_primary,
            key=lambda item: abs(item.section_index - secondary.section_index),
        )
        primary_pages = _section_pages(sections[primary.section_index])
        secondary_pages = _section_pages(sections[secondary.section_index])
        if primary_pages and secondary_pages:
            page_gap = min(
                abs(left - right)
                for left in primary_pages for right in secondary_pages
            )
            if page_gap > 2:
                continue
        elif abs(primary.section_index - secondary.section_index) > max_section_gap:
            continue
        primary_tokens = estimate_tokens(primary.abstract_text)
        secondary_tokens = estimate_tokens(secondary.abstract_text)
        size_similarity = min(primary_tokens, secondary_tokens) / max(
            primary_tokens, secondary_tokens,
        )
        if size_similarity < 0.45:
            continue
        selected_span = (
            _extend_secondary_abstract_frontmatter(
                sections[secondary.section_index].content, secondary,
            )
            if secondary.kind == "abstract"
            else secondary
        )
        selected.setdefault(secondary.section_index, []).append(selected_span)
    return primary_language, selected


def _expand_paper_secondary_frontmatter(
    sections: list[Section],
    selected: dict[int, list[_BilingualSummarySpan]],
) -> dict[int, list[_BilingualSummarySpan]]:
    """Recover translated front matter fragmented by PDF layout headings.

    Journal PDFs frequently split a translated title, byline and abstract into
    many synthetic sections, or continue the abstract on the next page.  Once
    an explicit, paired secondary abstract has been established, script and
    page continuity are sufficient to recover those fragments without using
    any document- or domain-specific vocabulary.
    """

    expanded = {
        index: list(values) for index, values in selected.items()
    }
    seeds = [
        span for values in selected.values() for span in values
        if span.kind == "abstract"
    ]
    for seed in seeds:
        section_index = seed.section_index
        language = seed.language
        seed_pages = _section_pages(sections[section_index])

        # Recover same-page translated title and byline fragments immediately
        # preceding the explicit abstract marker.  A partial tail terminates
        # the backwards scan so primary-language content is never crossed.
        for previous_index in range(section_index - 1, max(-1, section_index - 9), -1):
            previous = sections[previous_index]
            previous_pages = _section_pages(previous)
            if seed_pages and previous_pages and not set(seed_pages).intersection(previous_pages):
                break
            trailing = _trailing_language_span(previous.content, language)
            if trailing is None:
                break
            start, end = trailing
            expanded.setdefault(previous_index, []).append(
                _BilingualSummarySpan(
                    previous_index, start, end, language,
                    previous.content[start:end].strip(), "abstract",
                )
            )
            if start > 0:
                break

        keyword_pattern = (
            _EN_KEYWORD_MARKER if language == "en" else _ZH_KEYWORD_MARKER
        )
        if keyword_pattern.search(
            sections[section_index].content, seed.start, seed.end,
        ):
            continue

        # Recover a next-page abstract continuation.  Stop at the keyword
        # section (including translated funding metadata on that section), a
        # primary-language section, or a real body heading.
        for next_index in range(
            section_index + 1, min(len(sections), section_index + 7),
        ):
            current = sections[next_index]
            pages = _section_pages(current)
            if seed_pages and pages and min(pages) - max(seed_pages) > 2:
                break
            title = current.title.strip()
            if (
                not _SYNTHETIC_PDF_PAGE_TITLE.fullmatch(title)
                and _PAPER_MARKERS.search(title)
                and not _ABSTRACT_TITLE.fullmatch(title)
            ):
                break
            if not _language_dominates(current.content, language, minimum=12):
                break
            expanded.setdefault(next_index, []).append(
                _BilingualSummarySpan(
                    next_index, 0, len(current.content), language,
                    current.content.strip(), "abstract",
                )
            )
            if keyword_pattern.search(current.content):
                break

    normalized: dict[int, list[_BilingualSummarySpan]] = {}
    for section_index, spans in expanded.items():
        unique: dict[tuple[int, int, str, str], _BilingualSummarySpan] = {}
        for span in spans:
            unique[(span.start, span.end, span.language, span.kind)] = span
        normalized[section_index] = sorted(
            unique.values(), key=lambda item: (item.start, item.end),
        )
    return normalized


def _paper_secondary_mirror_abstracts(
    sections: list[Section],
) -> tuple[str | None, dict[int, list[_BilingualSummarySpan]]]:
    """Find translated front-matter abstracts that duplicate the main language.

    This deliberately requires an explicit pair, a clear document-body
    language and comparable abstract sizes.  A merely multilingual document or
    an unpaired foreign-language abstract remains graph eligible.
    """

    abstract_spans = [
        span
        for section_index, section in enumerate(sections)
        for span in _paper_abstract_spans_for_section(section, section_index)
    ]
    primary_language, selected = _select_secondary_mirror_summaries(
        sections, abstract_spans, frontmatter_sections=16,
        max_section_gap=12, frontmatter_page_limit=3,
    )
    return primary_language, _expand_paper_secondary_frontmatter(
        sections, selected,
    )


def _paper_adjacent_article_preview_spans(
    sections: list[Section],
) -> dict[int, list[tuple[int, int]]]:
    """Find a next-article preview interleaved with prior references.

    Two-column journal exports sometimes place the next article's front matter
    before the previous article's remaining bibliography in extracted text.
    Docling may preserve that preview and the bibliography continuation in one
    section, or split the translated front matter across two adjacent sections.
    Requiring an earlier abstract, an earlier reference run, a new abstract,
    keywords and DOI before the continuation makes this a high-confidence
    layout boundary instead of a topic-specific heuristic.
    """

    selected: dict[int, list[tuple[int, int]]] = {}
    seen_article_frontmatter = False
    seen_references = False
    pending_frontmatter: list[int] = []
    for section_index, section in enumerate(sections):
        abstract_spans = _paper_abstract_spans_for_section(
            section, section_index,
        )
        reference_spans = bibliography_reference_spans(section.content)
        if seen_article_frontmatter and seen_references and reference_spans:
            reference_start = min(start for start, _ in reference_spans)
            preview_indices = [*pending_frontmatter, section_index]
            preview = "\n".join(
                sections[index].content
                if index != section_index
                else section.content[:reference_start]
                for index in preview_indices
            )
            has_new_abstract = any(
                _paper_abstract_spans_for_section(sections[index], index)
                for index in pending_frontmatter
            ) or any(span.start < reference_start for span in abstract_spans)
            has_keywords = bool(
                _ZH_KEYWORD_MARKER.search(preview)
                or _EN_KEYWORD_MARKER.search(preview)
            )
            if (
                len(preview.strip()) >= 120
                and has_new_abstract
                and has_keywords
                and _DOI_MARKER.search(preview)
            ):
                for index in pending_frontmatter:
                    selected[index] = [(0, len(sections[index].content))]
                selected[section_index] = [(0, reference_start)]
            pending_frontmatter = []
        elif seen_article_frontmatter and seen_references:
            # A Chinese front-matter section is often followed by a translated
            # abstract whose tail resumes the previous article's references.
            # Hold only a structurally complete front-matter candidate and do
            # not classify it unless that continuation is actually observed.
            preview = section.content
            has_keywords = bool(
                _ZH_KEYWORD_MARKER.search(preview)
                or _EN_KEYWORD_MARKER.search(preview)
            )
            if (
                abstract_spans
                and has_keywords
                and _DOI_MARKER.search(preview)
            ):
                pending_frontmatter = [section_index]
            elif pending_frontmatter and abstract_spans and has_keywords:
                pending_frontmatter.append(section_index)
                pending_frontmatter = pending_frontmatter[-2:]
            else:
                pending_frontmatter = []
        if abstract_spans:
            seen_article_frontmatter = True
        if reference_spans or _REFERENCE_TITLE.fullmatch(section.title.strip()):
            seen_references = True
    return selected


def _paper_back_matter_spans(content: str) -> list[tuple[int, int]]:
    """Keep author biographies searchable without treating them as facts."""

    match = _AUTHOR_BIOGRAPHY_MARKER.search(content)
    return [(match.start(), len(content))] if match else []


def _paper_reference_continuation_spans(
    sections: list[Section],
) -> dict[int, list[tuple[int, int]]]:
    """Recover a citation entry tail split before the next page's first item."""

    selected: dict[int, list[tuple[int, int]]] = {}
    seen_reference_run = False
    for section_index, section in enumerate(sections):
        local_spans = bibliography_reference_spans(section.content)
        abstract_spans = _paper_abstract_spans_for_section(
            section, section_index,
        )
        if seen_reference_run and abstract_spans and not local_spans:
            # A structurally complete next article starts a new document unit;
            # its later numbered prose must not inherit the old bibliography.
            seen_reference_run = False
        if seen_reference_run and local_spans:
            first_start = min(start for start, _ in local_spans)
            if first_start > 0 and section.content[:first_start].strip():
                selected[section_index] = [(0, first_start)]
        if (
            local_spans
            or _REFERENCE_TITLE.fullmatch(section.title.strip())
            or re.search(r"(?m)^\s*参考\s*文献\s*$", section.content)
        ):
            seen_reference_run = True
    return selected


def _annotate_general_bilingual_mirror_summaries(
    sections: list[Section],
) -> list[Section]:
    """Mark explicit translated summaries for retrieval-only general chunks."""

    summary_spans = [
        span
        for section_index, section in enumerate(sections)
        for span in _paper_abstract_spans_for_section(section, section_index)
    ]
    summary_spans.extend(_general_summary_heading_spans(sections))
    primary_language, selected = _select_secondary_mirror_summaries(
        sections, summary_spans, frontmatter_sections=8,
    )
    if not selected:
        return sections

    annotated: list[Section] = []
    for section_index, section in enumerate(sections):
        spans = selected.get(section_index, [])
        if not spans:
            annotated.append(section)
            continue
        exclusion_spans = [
            {
                "start": span.start,
                "end": span.end,
                "reason": (
                    "secondary_language_mirror_abstract"
                    if span.kind == "abstract"
                    else "secondary_language_mirror_summary"
                ),
                "summary_kind": span.kind,
                "summary_language": span.language,
            }
            for span in sorted(spans, key=lambda item: item.start)
        ]
        annotated.append(Section(
            key=section.key,
            title=section.title,
            level=section.level,
            content=section.content,
            metadata={
                **section.metadata,
                "document_primary_language": primary_language,
                "graph_exclusion_spans": exclusion_spans,
                "bilingual_mirror_detection": "paired_structural_summary",
            },
        ))
    return annotated


def _paper_chunks(sections: list[Section], source_name: str) -> list[TextChunk]:
    """Build paper chunks with retrieval-only references and mirror abstracts.

    Many PDFs do not expose a real ``参考文献`` heading to the parser.  Detect
    only publication-shaped numbered lines, then isolate that contiguous run.
    Paired translated abstracts remain searchable but do not duplicate the
    primary-language facts in the graph.  Ordinary multilingual prose,
    numbered requirements and procedures remain graph eligible.
    """

    primary_language, secondary_abstracts = _paper_secondary_mirror_abstracts(
        sections,
    )
    adjacent_article_previews = _paper_adjacent_article_preview_spans(sections)
    reference_continuations = _paper_reference_continuation_spans(sections)

    def content_parts(
        content: str, *, explicit_reference: bool,
        explicit_abstract: bool, abstract_language: str | None,
        mirror_spans: list[_BilingualSummarySpan],
        adjacent_preview_spans: list[tuple[int, int]],
        back_matter_spans: list[tuple[int, int]],
        reference_continuation_spans: list[tuple[int, int]],
    ) -> list[tuple[str, str, str, str | None]]:
        if explicit_reference:
            return [("references", content, "explicit_heading", None)]
        reference_spans = tuple(sorted(
            (*bibliography_reference_spans(content),
             *reference_continuation_spans),
        ))
        if (
            not reference_spans
            and not mirror_spans
            and not adjacent_preview_spans
            and not back_matter_spans
        ):
            return [(
                "abstract" if explicit_abstract else "body",
                content,
                "explicit_heading" if explicit_abstract else "none",
                abstract_language if explicit_abstract else None,
            )]

        lines = content.splitlines(keepends=True)
        groups: list[tuple[str, str, str | None, list[str]]] = []
        offset = 0
        for line in lines:
            line_end = offset + len(line)
            mirror = next((
                span for span in mirror_spans
                if span.start < line_end and offset < span.end
            ), None)
            adjacent_preview = any(
                start < line_end and offset < end
                for start, end in adjacent_preview_spans
            )
            back_matter = any(
                start < line_end and offset < end
                for start, end in back_matter_spans
            )
            if back_matter:
                part, detection, language = (
                    "document_back_matter", "author_biography_heading", None,
                )
            elif any(
                start < line_end and offset < end
                for start, end in reference_spans
            ):
                part, detection, language = (
                    "references", "citation_structure", None,
                )
            elif adjacent_preview:
                part, detection, language = (
                    "adjacent_article_preview",
                    "interleaved_reference_page_boundary", None,
                )
            elif mirror is not None:
                part, detection, language = (
                    "secondary_abstract", "paired_bilingual_abstract",
                    mirror.language,
                )
            else:
                part = "abstract" if explicit_abstract else "body"
                detection = "explicit_heading" if explicit_abstract else "none"
                language = abstract_language if explicit_abstract else None
            if not groups or groups[-1][:3] != (part, detection, language):
                groups.append((part, detection, language, []))
            groups[-1][3].append(line)
            offset = line_end
        return [
            (part, "".join(values).strip(), detection, language)
            for part, detection, language, values in groups
            if "".join(values).strip()
        ] or [("body", content, "none", None)]

    chunks: list[TextChunk] = []
    for section_index, section in enumerate(sections):
        title = section.title.strip()
        is_reference = bool(_REFERENCE_TITLE.fullmatch(title))
        is_abstract = bool(_ABSTRACT_TITLE.fullmatch(title))
        explicit_abstract_language = (
            "zh" if is_abstract and "摘" in title
            else "en" if is_abstract else None
        )
        pieces: list[tuple[str, str, str, str | None]] = []
        for paper_part, content, detection, part_language in content_parts(
            section.content,
            explicit_reference=is_reference,
            explicit_abstract=is_abstract,
            abstract_language=explicit_abstract_language,
            mirror_spans=secondary_abstracts.get(section_index, []),
            adjacent_preview_spans=adjacent_article_previews.get(
                section_index, [],
            ),
            back_matter_spans=_paper_back_matter_spans(section.content),
            reference_continuation_spans=(
                [] if adjacent_article_previews.get(section_index)
                else reference_continuations.get(section_index, [])
            ),
        ):
            pieces.extend(
                (paper_part, body, detection, part_language)
                for body in _token_pieces(
                    content,
                    Config.paper_chunk_target_tokens,
                    Config.paper_chunk_max_tokens,
                    Config.paper_chunk_overlap_percent,
                )
            )
        for index, (paper_part, body, detection, part_language) in enumerate(
            pieces, start=1,
        ):
            graph_enabled = (
                paper_part not in {
                    "references", "secondary_abstract",
                    "adjacent_article_preview", "document_back_matter",
                }
                and bool(section.metadata.get("graph_extraction_enabled", True))
            )
            metadata: dict[str, Any] = {
                "paper_part": paper_part,
                "chunk_role": "child",
                "reference_detection": detection,
                "graph_extraction_enabled": graph_enabled,
            }
            if primary_language:
                metadata["document_primary_language"] = primary_language
            if part_language:
                metadata["abstract_language"] = part_language
            if paper_part == "secondary_abstract":
                metadata["graph_exclusion_reason"] = (
                    "secondary_language_mirror_abstract"
                )
                metadata["mirror_abstract_of_language"] = primary_language
            elif paper_part == "adjacent_article_preview":
                metadata["graph_exclusion_reason"] = (
                    "adjacent_article_preview_interleaved_with_references"
                )
            elif paper_part == "document_back_matter":
                metadata["graph_exclusion_reason"] = (
                    "author_biography_document_back_matter"
                )
            chunks.append(_chunk(
                section, source_name, body, index, len(pieces), "paper",
                metadata=metadata,
            ))
    return chunks


def _table_chunks(blocks: list[DocumentBlock], sections: list[Section], source_name: str) -> list[TextChunk]:
    section_map = {item.key: item for item in sections}
    rows = [block for block in blocks if block.block_type == "table_row"]
    if not rows:
        raise ValueError("表格分块策略未发现可用表格行，请改用通用策略或选择含表格的文档")
    counts: dict[str, int] = {}
    for row in rows:
        counts[row.section_key] = counts.get(row.section_key, 0) + 1
    positions: dict[str, int] = {}
    chunks: list[TextChunk] = []
    for row in rows:
        positions[row.section_key] = positions.get(row.section_key, 0) + 1
        section = section_map.get(row.section_key) or Section(row.section_key, row.section_title, row.level, row.text)
        metadata = {"chunk_role": "table_row", "table": row.metadata,
                    "graph_extraction_enabled": True}
        chunks.append(_chunk(section, source_name, row.text, positions[row.section_key],
                             counts[row.section_key], "table", key_suffix=f"row{row.order:04d}",
                             metadata=metadata))
    return chunks


def parse_document_structure(path: Path, *, requested_parser: str = "auto",
                             requested_strategy: str = "auto") -> tuple[ProcessingPlan, ParsedDocument]:
    """Route and parse a document without creating any retrieval chunks."""
    plan = route_processing(path, requested_parser, requested_strategy)
    try:
        parsed = parse_document(path, plan.effective_parser)
    except Exception as exc:
        # Docling's optional layout/OCR models may be unavailable in an
        # offline deployment even though the PDF has a usable text layer.
        # Automatic routing must keep the deterministic text path available;
        # an explicitly requested Docling parse remains fail-closed.
        if (
            plan.effective_parser == "docling"
            and requested_parser == "auto"
            and plan.profile.has_text_layer
        ):
            LOGGER.warning(
                "Docling unavailable for %s; using the available PDF text layer: %s",
                path.name, exc,
            )
            parsed = parse_document(path, "plaintext_pdf")
            plan = ProcessingPlan(
                plan.requested_parser,
                plan.requested_strategy,
                "plaintext_pdf",
                plan.effective_strategy,
                plan.routing_reason + "；Docling 版面模型不可用，回退 PDF 文本层解析",
                plan.profile,
            )
        elif not (
            plan.effective_parser == "plaintext_pdf"
            and requested_parser == "auto"
            and docling_available()
        ):
            raise
        else:
            parsed = parse_document(path, "docling")
            plan = ProcessingPlan(
                plan.requested_parser,
                plan.requested_strategy,
                "docling",
                plan.effective_strategy,
                plan.routing_reason + "；轻量 PDF 解析失败，回退 Docling",
                plan.profile,
            )
    if (plan.effective_parser == "plaintext_pdf" and requested_parser == "auto"
            and parsed.quality.get("text_chars", 0) < max(200, plan.profile.page_count * 40)
            and docling_available()):
        try:
            parsed = parse_document(path, "docling")
            plan = ProcessingPlan(
                plan.requested_parser,
                plan.requested_strategy,
                "docling",
                plan.effective_strategy,
                plan.routing_reason + "；PDF 文本质量不足，回退 Docling",
                plan.profile,
            )
        except Exception as exc:
            if not plan.profile.has_text_layer:
                raise
            LOGGER.warning(
                "Docling quality fallback unavailable for %s; keeping PDF text layer: %s",
                path.name, exc,
            )
    return plan, sanitize_parsed_document(parsed)


def build_processing_result(path: Path, plan: ProcessingPlan, parsed: ParsedDocument, *,
                            source_name: str = "", chunk_max_tokens: int | None = None) -> ProcessingResult:
    """Build sections and chunks from an already persisted structure parse."""
    parsed = sanitize_parsed_document(parsed)
    sections = blocks_to_sections(parsed.blocks)
    if not sections:
        raise ValueError("未提取到可入库章节")
    strategy = plan.effective_strategy
    if strategy in {"default", "general"}:
        sections = _annotate_general_bilingual_mirror_summaries(sections)
    if strategy == "default":
        chunks = split_sections(sections, max_chars=Config.chunk_max_chars,
                                overlap_chars=Config.chunk_overlap_chars, source_name=source_name or path.name)
    elif strategy == "general":
        sections = _annotate_general_structure_roles(sections)
        sections = _coalesce_general_micro_sections(
            sections, target_tokens=Config.general_chunk_target_tokens,
        )
        chunks = _general_chunks(sections, source_name or path.name,
                                 target=Config.general_chunk_target_tokens,
                                 maximum=Config.general_chunk_max_tokens,
                                 overlap=Config.general_chunk_overlap_percent)
    elif strategy == "paper":
        chunks = _paper_chunks(sections, source_name or path.name)
    else:
        chunks = _table_chunks(parsed.blocks, sections, source_name or path.name)
    if chunk_max_tokens is not None:
        chunks = _enforce_chunk_token_limit(chunks, int(chunk_max_tokens))
    if not chunks:
        raise ValueError("分块完成但没有生成知识块")
    return ProcessingResult(plan, parsed, sections, chunks)


def process_document(path: Path, *, requested_parser: str = "auto", requested_strategy: str = "auto",
                     source_name: str = "", chunk_max_tokens: int | None = None) -> ProcessingResult:
    plan, parsed = parse_document_structure(
        path,
        requested_parser=requested_parser,
        requested_strategy=requested_strategy,
    )
    return build_processing_result(
        path,
        plan,
        parsed,
        source_name=source_name,
        chunk_max_tokens=chunk_max_tokens,
    )
