import pytest
from docx import Document as DocxDocument

import knowledge_forge.chunking as chunking
from knowledge_forge.chunking import (
    DocumentBlock,
    DocumentProfile,
    ParsedDocument,
    ProcessingPlan,
    Section,
    build_processing_result,
    extract_plain_sections,
    parse_document_structure,
    process_document,
    route_processing,
    split_sections,
    validate_processing_request,
)


def test_general_sentence_segmentation_uses_shared_exact_spans():
    from knowledge_forge.chunking import _segment_spans

    content = (
        "Version v1.2.3 costs 3.14 USD. Visit https://example.com/a.b and a.b@example.com. "
        "Dr. Smith said ready... 然后继续……完成！分块: 正文。"
    )
    spans = _segment_spans(content)
    assert "".join(content[start:end] for start, end, _ in spans) == content
    assert all(content[start:end] == value for start, end, value in spans)
    assert [value for _, _, value in spans] == [
        "Version v1.2.3 costs 3.14 USD. ",
        "Visit https://example.com/a.b and a.b@example.com. ",
        "Dr. Smith said ready... ",
        "然后继续……",
        "完成！",
        "分块: 正文。",
    ]


def test_docling_page_break_markers_preserve_block_page_provenance():
    marker = chunking._DOCLING_PAGE_BREAK_MARKER
    blocks = chunking._text_to_blocks(
        f"# 第一章\n\n第一页正文。\n\n{marker}\n\n第二页正文。",
        "测试文档",
        page=1,
        page_break_marker=marker,
    )

    paragraphs = [block for block in blocks if block.block_type == "paragraph"]
    assert [(block.text, block.page) for block in paragraphs] == [
        ("第一页正文。", 1),
        ("第二页正文。", 2),
    ]
    sections = chunking.blocks_to_sections(blocks)
    assert sections[0].metadata["pages"] == [1, 2]


def test_docling_layout_value_and_sentence_fragments_are_not_headings():
    blocks = chunking._text_to_blocks(
        "## 3.7.2 开销对比\n\n"
        "正文。\n\n"
        "## 41.84 ms Dropout\n\n"
        "较基线节省时间。\n\n"
        "## 150.39 MB\n\n"
        "显存占用基本持平。\n\n"
        "## 持续低迷。\n\n"
        "后续说明。\n\n"
        "## 4 结论\n\n"
        "结论正文。",
        "paper",
        page=12,
        validate_layout_headings=True,
    )

    assert [item.text for item in blocks if item.block_type == "heading"] == [
        "3.7.2 开销对比",
        "4 结论",
    ]
    paragraphs = [item.text for item in blocks if item.block_type == "paragraph"]
    assert any("41.84 ms Dropout" in value for value in paragraphs)
    assert any("150.39 MB" in value for value in paragraphs)
    assert any("持续低迷。" in value for value in paragraphs)


def test_authored_markdown_measurement_heading_is_preserved():
    blocks = chunking._text_to_blocks(
        "## 150.39 MB\n\n作者明确定义的标题。",
        "manual",
    )

    assert blocks[0].block_type == "heading"
    assert blocks[0].text == "150.39 MB"


def test_docling_unmarked_measurement_is_not_a_numbered_heading():
    blocks = chunking._text_to_blocks(
        "3.7.2 开销对比\n\n"
        "41.84 ms Dropout\n\n"
        "150.39 MB\n\n"
        "4 结论\n\n"
        "结论正文。",
        "paper",
        page=12,
        validate_layout_headings=True,
    )

    assert [item.text for item in blocks if item.block_type == "heading"] == [
        "3.7.2 开销对比",
        "4 结论",
    ]
    paragraphs = [item.text for item in blocks if item.block_type == "paragraph"]
    assert any("41.84 ms Dropout" in value for value in paragraphs)
    assert any("150.39 MB" in value for value in paragraphs)


def test_markdown_multilevel_table_header_is_not_emitted_as_data_row(tmp_path):
    source = tmp_path / "multi-header.md"
    source.write_text(
        "| Methods | Dice/% (↑) | Dice/% (↑) | HD95(↓) | HD95(↓) |\n"
        "| --- | --- | --- | --- | --- |\n"
        "| Methods | C1 | C2 | C1 | C2 |\n"
        "| Baseline | Atlas | Mapping | Survey | Analysis |\n",
        encoding="utf-8",
    )

    _plan, parsed = parse_document_structure(source)
    rows = [block for block in parsed.blocks if block.block_type == "table_row"]

    assert len(rows) == 1
    assert rows[0].text == (
        "Methods: Baseline；Dice/% (↑) / C1: Atlas；"
        "Dice/% (↑) / C2: Mapping；HD95(↓) / C1: Survey；"
        "HD95(↓) / C2: Analysis"
    )


def test_structure_parsing_can_be_inspected_before_chunking(tmp_path):
    source = tmp_path / "staged.md"
    source.write_text("# 第一章\n\n正文内容。", encoding="utf-8")

    plan, parsed = parse_document_structure(source)

    assert plan.effective_parser == "native_text"
    assert [block.block_type for block in parsed.blocks] == ["heading", "paragraph"]

    result = build_processing_result(source, plan, parsed, source_name=source.name)
    assert result.chunks
    assert result.chunks[0].section_key == parsed.blocks[1].section_key


def test_auto_docling_failure_falls_back_to_available_pdf_text_layer(monkeypatch, tmp_path):
    source = tmp_path / "complex-manual.pdf"
    source.write_bytes(b"placeholder")
    profile = DocumentProfile(
        suffix=".pdf",
        page_count=2,
        text_chars=500,
        has_text_layer=True,
        layout_complexity="complex",
    )
    plan = ProcessingPlan(
        "auto", "auto", "docling", "general", "复杂版面", profile,
    )
    parsed = ParsedDocument(
        [DocumentBlock("b1", "paragraph", "正文", 1, "s1", "说明")],
        "plaintext_pdf", "pypdf-test", {"text_chars": 500},
    )
    calls = []

    monkeypatch.setattr(chunking, "route_processing", lambda *args, **kwargs: plan)

    def parse(path, parser):
        calls.append(parser)
        if parser == "docling":
            raise RuntimeError("offline model unavailable")
        return parsed

    monkeypatch.setattr(chunking, "parse_document", parse)
    result_plan, result = chunking.parse_document_structure(source)

    assert calls == ["docling", "plaintext_pdf"]
    assert result_plan.effective_parser == "plaintext_pdf"
    assert "回退 PDF 文本层解析" in result_plan.routing_reason
    assert result is parsed


def test_split_keeps_heading_metadata():
    chunks = split_sections(
        [Section("s1", "第一章", 1, "甲。" * 100)],
        max_chars=180,
        overlap_chars=20,
        source_name="manual.txt",
    )

    assert len(chunks) > 1
    assert all(chunk.title == "第一章" for chunk in chunks)
    assert all("来源文档: manual.txt" in chunk.content for chunk in chunks)


def test_markdown_heading_sections(tmp_path):
    source = tmp_path / "sample.md"
    source.write_text("# 总则\n正文\n\n## 子节\n内容", encoding="utf-8")

    sections = extract_plain_sections(source)

    assert [item.title for item in sections] == ["总则", "子节"]


def test_native_text_blocks_preserve_nested_list_indentation():
    blocks = chunking._text_to_blocks(
        "3 核心系统\n\n"
        "— 分区管理\n"
        "  — 空间隔离\n"
        "— 任务管理\n",
        "manual",
    )

    paragraph = next(
        item for item in blocks if item.block_type == "paragraph"
    )
    assert paragraph.text.splitlines() == [
        "— 分区管理",
        "  — 空间隔离",
        "— 任务管理",
    ]


def test_general_chunks_preserve_nested_list_indentation(tmp_path):
    source = tmp_path / "nested-list.txt"
    source.write_text(
        "3 核心系统\n\n"
        "— 分区管理\n"
        "  — 空间隔离\n"
        "  — 分区间通信\n"
        "— 任务管理\n"
        "  — 调度表管理\n",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="general", source_name=source.name,
    )

    content = "\n".join(chunk.content for chunk in result.chunks)
    assert "— 分区管理\n  — 空间隔离\n  — 分区间通信" in content
    assert "— 任务管理\n  — 调度表管理" in content


def test_plaintext_pdf_style_table_rows_are_not_numbered_headings():
    blocks = chunking._text_to_blocks(
        "3.1 故障码定义\n\n"
        "序号   层次     功能模块       故障原因       故障码\n"
        "1      CORE     自旋锁         初始化错误     ACE_ERROR_CODE\n"
        "2      CORE     事件管理       状态错误       ACE_STATE_ERROR\n",
        "manual",
        page=3,
    )

    assert blocks[0].block_type == "heading"
    assert blocks[0].text == "3.1 故障码定义"
    assert [item.block_type for item in blocks[1:]] == ["paragraph"]
    assert "1      CORE" in blocks[1].text
    assert "2      CORE" in blocks[1].text

    compact = chunking._text_to_blocks(
        "3.2 故障信息表\n"
        "1 CORE 自旋锁 CPC 为直接处理模式\n"
        "2 ACE_ERROR_CODE_INT_ERROR 中断错误\n",
        "manual",
        page=3,
    )
    assert compact[0].block_type == "heading"
    assert [item.block_type for item in compact[1:]] == ["paragraph"]

    legitimate_api_heading = chunking._text_to_blocks(
        "3 API Reference\n", "manual",
    )
    assert legitimate_api_heading[0].block_type == "heading"


def test_pdf_versions_dates_and_numbered_api_rows_are_not_headings():
    blocks = chunking._text_to_blocks(
        "0.0.0.9/ACoreIDE\n"
        "2023 年 8 月\n"
        "9 RESUME\n"
        "12 START\n"
        "21 REPLENISH\n"
        "62 分区间通信（队\n"
        "197 atexit() Y Y N 需要进程\n"
        "198 clock() Y Y N 需要进程\n",
        "manual",
        page=5,
    )

    assert [item.block_type for item in blocks] == ["paragraph"]
    assert "0.0.0.9/ACoreIDE" in blocks[0].text
    assert "197 atexit() Y Y N" in blocks[0].text


def test_legitimate_numbered_headings_remain_supported():
    blocks = chunking._text_to_blocks(
        "1. 概述\n正文\n"
        "2.3 API Reference\n接口说明\n"
        "3.1. Advanced Usage\n高级用法\n",
        "manual",
    )

    assert [item.text for item in blocks if item.block_type == "heading"] == [
        "1. 概述",
        "2.3 API Reference",
        "3.1. Advanced Usage",
    ]


def test_pdf_decimal_cells_are_not_headings_but_spaced_headings_are():
    blocks = chunking._text_to_blocks(
        "3 . 3 实验设置\n"
        "87 . 10\n"
        "HD 95 ( ↓ )\n"
        "C 1\n",
        "manual",
        page=8,
    )

    assert blocks[0].block_type == "heading"
    assert blocks[0].text == "3 . 3 实验设置"
    assert [item.block_type for item in blocks[1:]] == ["paragraph"]
    assert "87 . 10\nHD 95 ( ↓ )\nC 1" in blocks[1].text


def test_table_of_contents_dot_leader_is_not_a_heading():
    blocks = chunking._text_to_blocks(
        "1 概述 ................................................ 4\n",
        "manual",
    )

    assert len(blocks) == 1
    assert blocks[0].block_type == "paragraph"


def test_incompatible_manual_parser_is_rejected_before_ingestion():
    with pytest.raises(ValueError, match="不能处理"):
        validate_processing_request("manual.docx", "spreadsheet", "general")


def test_csv_auto_routes_to_spreadsheet_table_chunks(tmp_path):
    source = tmp_path / "industry.csv"
    source.write_text("企业,产品,地区\n甲公司,芯片,上海\n乙公司,设备,苏州\n", encoding="utf-8")

    result = process_document(source, source_name=source.name)

    assert result.plan.effective_parser == "spreadsheet"
    assert result.plan.effective_strategy == "table"
    assert len(result.chunks) == 2
    assert "企业: 甲公司" in result.chunks[0].content
    assert result.chunks[0].metadata["chunk_role"] == "table_row"


def test_xlsx_preserves_sheet_headers_in_each_row_chunk(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "catalog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品目录"
    sheet.append(["型号", "名称", "状态"])
    sheet.append(["A-01", "控制器", "在产"])
    workbook.save(source)

    result = process_document(source, source_name=source.name)

    assert result.plan.effective_parser == "spreadsheet"
    assert len(result.chunks) == 1
    assert "型号: A-01" in result.chunks[0].content
    assert result.chunks[0].metadata["table"]["sheet"] == "产品目录"


def test_general_strategy_builds_parent_and_child_chunks(tmp_path):
    source = tmp_path / "manual.md"
    source.write_text("# 操作说明\n\n" + "这是一个较长的操作步骤。" * 500, encoding="utf-8")

    result = process_document(source, requested_strategy="general", source_name=source.name)

    parents = [item for item in result.chunks if item.metadata.get("chunk_role") == "parent"]
    children = [item for item in result.chunks if item.metadata.get("chunk_role") == "child"]
    assert len(parents) == 1
    assert len(children) > 1
    assert all(item.parent_key == parents[0].key for item in children)
    assert parents[0].metadata["graph_extraction_enabled"] is False


def test_general_strategy_coalesces_many_micro_sections_without_loss(tmp_path):
    source = tmp_path / "api-reference.md"
    source.write_text("\n".join(
        f"## 2.{index} 接口{index}\n接口{index}用于处理参数{index}。"
        for index in range(1, 31)
    ), encoding="utf-8")

    result = process_document(
        source, requested_strategy="general", source_name=source.name,
    )

    assert len(result.sections) < 30
    assert len(result.chunks) < 30
    combined = "\n".join(item.content for item in result.chunks)
    assert "接口1用于处理参数1" in combined
    assert "接口30用于处理参数30" in combined
    bundled = [
        item for item in result.sections
        if item.metadata.get("coalesced_section_count")
    ]
    assert bundled
    assert sum(
        int(item.metadata["coalesced_section_count"]) for item in bundled
    ) == 30
    assert all(item.metadata.get("source_section_keys") for item in bundled)


def test_general_strategy_coalesces_synthetic_pdf_page_sections(monkeypatch, tmp_path):
    source = tmp_path / "manual.pdf"
    source.write_bytes(b"placeholder")
    sections = [
        Section(
            f"p{index:04d}", f"第 {index} 页", 1,
            f"第 {index} 页上的连续正文。",
            {"pages": [index], "block_types": ["paragraph"]},
        )
        for index in range(1, 31)
    ]
    parsed = ParsedDocument(
        [
            DocumentBlock(
                f"b{index:05d}", "paragraph", section.content, index,
                section.key, section.title, section.level, (), index,
            )
            for index, section in enumerate(sections, start=1)
        ],
        "plaintext_pdf", "pypdf-test", {"text_chars": 600},
    )
    plan = ProcessingPlan(
        "auto", "general", "plaintext_pdf", "general", "test",
        DocumentProfile(suffix=".pdf", page_count=30, text_chars=600),
    )

    result = build_processing_result(
        source, plan, parsed, source_name=source.name,
    )

    assert len(result.sections) < len(sections)
    assert sum(
        int(item.metadata.get("coalesced_section_count") or 0)
        for item in result.sections
    ) == len(sections)
    combined = "\n".join(item.content for item in result.chunks)
    assert "第 1 页上的连续正文" in combined
    assert "第 30 页上的连续正文" in combined


def test_general_toc_remains_searchable_but_skips_graph_extraction(tmp_path):
    source = tmp_path / "manual.md"
    source.write_text(
        "# 目录\n"
        "1 概述 ................................ 1\n"
        "2 安装 ................................ 3\n"
        "3 配置 ................................ 8\n"
        "4 故障处理 ........................... 12\n",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="general", source_name=source.name,
    )

    assert len(result.sections) == 1
    assert result.sections[0].metadata["structural_role"] == "table_of_contents"
    assert result.sections[0].metadata["toc_line_count"] == 4
    assert result.chunks
    assert "故障处理" in "\n".join(chunk.content for chunk in result.chunks)
    assert all(
        chunk.metadata["graph_extraction_enabled"] is False
        for chunk in result.chunks
    )


def test_general_toc_gate_keeps_prose_and_plain_feature_lists_graph_eligible(tmp_path):
    source = tmp_path / "features.md"
    source.write_text(
        "# 核心操作系统\n"
        "本节说明系统提供的能力及其约束。\n"
        "另见安装步骤 ........................ 3\n"
        "另见配置步骤 ........................ 8\n"
        "— 分区管理\n"
        "— 任务管理\n"
        "— 时间管理\n"
        "这些能力共同构成核心服务。\n",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="general", source_name=source.name,
    )

    assert result.sections[0].metadata.get("structural_role") is None
    assert all(
        chunk.metadata["graph_extraction_enabled"] is True
        for chunk in result.chunks
        if chunk.metadata.get("chunk_role") == "child"
    )


def test_project_chunk_token_limit_is_a_hard_bound(tmp_path):
    source = tmp_path / "bounded.md"
    source.write_text("# 说明\n\n" + "知识图谱需要稳定的项目级分块上限。" * 300, encoding="utf-8")

    result = process_document(
        source,
        requested_strategy="general",
        source_name=source.name,
        chunk_max_tokens=128,
    )

    assert len(result.chunks) > 1
    assert all(chunk.token_count <= 128 for chunk in result.chunks)
    assert any(chunk.metadata.get("token_limit_split") for chunk in result.chunks)


def test_paper_references_remain_searchable_but_skip_graph_extraction(tmp_path):
    source = tmp_path / "paper.md"
    source.write_text(
        "# 摘要\n研究内容。\n# 引言\n背景。\n# 研究方法\n方法。\n"
        "# 实验结果\n结果。\n# 参考文献\n[1] Example, 2026.",
        encoding="utf-8",
    )

    plan = route_processing(source)
    result = process_document(source, requested_strategy="paper", source_name=source.name)

    assert plan.effective_strategy == "paper"
    references = [item for item in result.chunks if item.metadata.get("paper_part") == "references"]
    assert references
    assert references[0].metadata["graph_extraction_enabled"] is False


def test_paper_keeps_paired_translation_abstract_searchable_but_out_of_graph(tmp_path):
    source = tmp_path / "bilingual-paper.md"
    source.write_text(
        "# 第 1 页\n"
        "【摘要】联邦学习在不共享原始数据的情况下协同训练模型，并通过本地更新保护隐私。"
        "本文总结其方法、应用、限制与发展方向。\n"
        "【关键词】联邦学习；隐私保护；医学数据\n"
        "Federated learning for privacy-preserving medical collaboration\n"
        "【Abstract】Federated learning collaboratively trains models without sharing raw data "
        "and protects privacy through local updates. This paper summarizes its methods, "
        "applications, limitations, and future directions.\n"
        "【Keywords】Federated learning; privacy; medical data\n"
        "# 引言\n" +
        "中文正文系统介绍研究背景、方法原理、应用场景、实验观察和局限性。" * 45,
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )
    translated = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "secondary_abstract"
    ]

    assert translated
    assert all(
        item.metadata["graph_extraction_enabled"] is False
        for item in translated
    )
    assert all(
        item.metadata["graph_exclusion_reason"]
        == "secondary_language_mirror_abstract"
        for item in translated
    )
    assert translated[0].metadata["document_primary_language"] == "zh"
    assert "Federated learning" in "\n".join(item.content for item in translated)
    assert any(
        "【摘要】" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_paper_recovers_fragmented_ocr_spaced_secondary_frontmatter(tmp_path):
    source = tmp_path / "fragmented-bilingual-paper.md"
    source.write_text(
        "# 第 1 页\n"
        "期刊页眉与中文论文标题。\n"
        "# 1，2\n"
        "摘 要：系统协调多个客户端，在不共享原始数据的条件下完成训练。"
        "本文介绍方法、实验、结果、局限与未来方向。\n"
        "关键词：系统；训练；验证\n"
        "基金项目：通用研究计划\n"
        "Personalized collaborative system for verified training\n"
        "DU Example\n"
        "# 1,2\n, WEI Zheng\n"
        "# 1,2\n, LEI Tao\n"
        "# 1,2\n"
        "(1. School of Engineering, Example University, China)\n"
        "Abstrac t: The system coordinates multiple clients without sharing raw data. "
        "This paper explains the method, experiments, results, limitations, and future work.\n"
        "# 第 2 页\n"
        "The objective function continues on the next page and the evaluation confirms "
        "stable collaborative training under heterogeneous conditions.\n"
        "Key words: system; training; verification\n"
        "Foundation Item: Generic research program\n"
        "# 0 引言\n" +
        "中文正文详细介绍架构、协议、实验设计、结果分析、运行权衡和适用限制。" * 45,
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )
    secondary = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "secondary_abstract"
    ]
    secondary_text = "\n".join(item.content for item in secondary)

    assert secondary
    assert all(
        item.metadata["graph_extraction_enabled"] is False
        for item in secondary
    )
    assert "Personalized collaborative system" in secondary_text
    assert "WEI Zheng" in secondary_text
    assert "Abstrac t" in secondary_text
    assert "objective function continues" in secondary_text
    assert "Foundation Item" in secondary_text
    assert any(
        "系统协调多个客户端" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )
    assert any(
        "中文正文详细介绍架构" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_paper_handles_chinese_translation_of_english_primary_abstract(tmp_path):
    source = tmp_path / "english-primary-paper.md"
    source.write_text(
        "# Abstract\n"
        "A distributed system coordinates independent nodes, preserves local state, and "
        "reports verified results. The paper describes its design, evaluation, and limits.\n"
        "# 摘要\n"
        "分布式系统协调独立节点、保留本地状态并报告经过验证的结果。"
        "本文介绍系统设计、评估方法及其局限。\n"
        "# Introduction\n" +
        "The main body explains architecture, protocols, experiments, operational tradeoffs, "
        "and observed limitations in a production-independent setting. " * 35,
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )
    chinese_abstract = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "secondary_abstract"
    ]

    assert chinese_abstract
    assert chinese_abstract[0].metadata["abstract_language"] == "zh"
    assert chinese_abstract[0].metadata["document_primary_language"] == "en"
    assert chinese_abstract[0].metadata["graph_extraction_enabled"] is False
    assert any(
        item.title == "Abstract"
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_paper_does_not_suppress_abstracts_when_body_is_truly_bilingual(tmp_path):
    source = tmp_path / "mixed-language-paper.md"
    source.write_text(
        "# 摘要\n"
        "系统协调多个节点并保留本地状态，本文说明系统方法、结果、限制与后续方向。\n"
        "# Abstract\n"
        "The system coordinates multiple nodes while retaining local state. The paper "
        "explains the method, results, limits, and future directions.\n"
        "# 中文正文\n" + "中文内容介绍方法、结果与限制。" * 18 + "\n"
        "# English body\n" +
        "English content explains methods, results, and limitations. " * 28,
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )

    assert not any(
        item.metadata.get("paper_part") == "secondary_abstract"
        for item in result.chunks
    )
    assert all(
        item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_docx_paper_uses_the_same_bilingual_abstract_isolation(tmp_path):
    source = tmp_path / "bilingual-paper.docx"
    document = DocxDocument()
    document.add_heading("摘要", level=1)
    document.add_paragraph(
        "系统协调多个节点、保留本地状态并输出验证结果。本文介绍设计、评估和限制。"
    )
    document.add_heading("Abstract", level=1)
    document.add_paragraph(
        "The system coordinates multiple nodes, retains local state, and produces verified "
        "results. This paper describes the design, evaluation, and limitations."
    )
    document.add_heading("正文", level=1)
    document.add_paragraph(
        "中文正文详细说明架构、协议、实验、运行权衡、观察结果和适用限制。" * 38
    )
    document.save(source)

    result = process_document(
        source,
        requested_parser="native_docx",
        requested_strategy="paper",
        source_name=source.name,
    )

    translated = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "secondary_abstract"
    ]
    assert translated
    assert translated[0].metadata["abstract_language"] == "en"
    assert translated[0].metadata["graph_extraction_enabled"] is False


def test_plain_text_paper_uses_the_same_bilingual_abstract_isolation(tmp_path):
    source = tmp_path / "bilingual-paper.txt"
    source.write_text(
        "【摘要】系统协调多个节点、保留本地状态并输出验证结果。"
        "本文介绍设计、评估和限制。\n"
        "【关键词】系统；验证；节点\n"
        "【Abstract】The system coordinates multiple nodes, retains local state, and produces "
        "verified results. This paper describes the design, evaluation, and limitations.\n"
        "【Keywords】system; verification; nodes\n" +
        "中文正文详细说明架构、协议、实验、运行权衡、观察结果和适用限制。" * 38,
        encoding="utf-8",
    )

    result = process_document(
        source,
        requested_parser="native_text",
        requested_strategy="paper",
        source_name=source.name,
    )

    translated = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "secondary_abstract"
    ]
    assert translated
    assert translated[0].metadata["abstract_language"] == "en"
    assert translated[0].metadata["graph_extraction_enabled"] is False


@pytest.mark.parametrize("strategy", ["general", "default"])
def test_ordinary_documents_isolate_only_explicit_paired_summary(
    tmp_path, strategy,
):
    source = tmp_path / f"ordinary-{strategy}.md"
    source.write_text(
        "# 执行摘要\n"
        "系统通过多个独立节点完成协作，保留本地状态并输出验证结果。"
        "本节概括设计、评估、运行边界和后续方向。\n"
        "# Executive Summary\n"
        "The system collaborates through multiple independent nodes, retains local state, "
        "and produces verified results. This section summarizes the design, evaluation, "
        "operational boundaries, and future directions.\n"
        "# 第一章\n" +
        "中文正文详细介绍架构设计、通信协议、实验方法、结果分析、运行权衡与适用限制。" * 42 +
        "\n# English Notes\n"
        "RFC 9000 defines a protocol requirement that is unique and must remain extractable.\n"
        "# Summary\n"
        "This unpaired operational summary contains a unique deployment fact and must remain "
        "eligible for graph extraction.",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy=strategy, source_name=source.name,
    )
    translated_summary = [
        item for item in result.chunks if item.title == "Executive Summary"
    ]

    assert translated_summary
    assert all(
        item.metadata["graph_extraction_enabled"] is False
        for item in translated_summary
    )
    assert translated_summary[0].metadata["summary_language"] == "en"
    assert translated_summary[0].metadata["document_primary_language"] == "zh"
    assert any(
        item.title == "执行摘要"
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )
    assert any(
        item.title == "English Notes"
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )
    assert any(
        item.title == "Summary"
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_general_document_partitions_inline_paired_abstract_without_losing_text(tmp_path):
    source = tmp_path / "inline-summary.txt"
    source.write_text(
        "【摘要】系统协调多个节点并保留本地状态，本文概括设计、评估、限制与后续方向。\n"
        "【关键词】系统；节点；验证\n"
        "【Abstract】The system coordinates multiple nodes and retains local state. This text "
        "summarizes the design, evaluation, limitations, and future directions.\n"
        "【Keywords】system; nodes; verification\n" +
        "中文正文详细介绍架构、协议、实验、结果、运行权衡与适用限制。" * 45,
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="general", source_name=source.name,
    )
    retrieval_only = [
        item for item in result.chunks
        if item.metadata.get("graph_exclusion_reason")
        == "secondary_language_mirror_abstract"
    ]

    assert retrieval_only
    assert all(
        item.metadata["graph_extraction_enabled"] is False
        for item in retrieval_only
    )
    assert "【Abstract】" in "\n".join(item.content for item in retrieval_only)
    assert any(
        "中文正文" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_paper_detects_unheaded_reference_run_without_hiding_numbered_steps(tmp_path):
    source = tmp_path / "mixed-paper.md"
    source.write_text(
        "# 结果\n"
        "[1] Scheduler initializes the timer before dispatch.\n"
        "The measured result remains available to the graph.\n"
        "[12] Zhang Q. A method[J]. Journal, 2024, 29(1): 10-18.\n"
        "[13] Li R. Another method[J]. Journal, 2025, 30(2): 20-28.\n",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )
    references = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "references"
    ]
    bodies = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "body"
    ]

    assert references
    assert all(
        item.metadata["graph_extraction_enabled"] is False
        for item in references
    )
    assert any("Zhang Q" in item.content for item in references)
    assert any("Scheduler initializes" in item.content for item in bodies)
    assert all(
        item.metadata["graph_extraction_enabled"] is True for item in bodies
    )


def test_paper_detects_fullwidth_cross_page_references_and_author_biography(tmp_path):
    source = tmp_path / "cross-page-references.md"
    source.write_text(
        "# 4 结论\n"
        "本文总结系统设计、实验结果与适用边界。\n"
        "参考文献\n"
        "［ 1 ］ Zhang Q. A verified method [ J ]. Journal, 2024, 10(1): 1-8.\n"
        "# 第 15 页\n"
        "continued citation text [ C ] // Conference, 2023: 9-12.\n"
        "［ 2 ］ Li R. Another method [ J ]. Journal, 2025, 11(2): 13-20.\n"
        "［ 3 ］ Wang T. Final method [ J ]. Journal, 2026, 12(3): 21-30.\n"
        "# 第 16 页\n"
        "［ 4 ］ Zhao M. Last method [ J ]. Journal, 2022, 13(4): 31-40.\n"
        "作 者 简 介：\n"
        "张某现为某机构研究人员，主要研究方向为系统设计。\n",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )
    references = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "references"
    ]
    biographies = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "document_back_matter"
    ]

    assert references
    assert "continued citation text" in "\n".join(
        item.content for item in references
    )
    assert "［ 2 ］" in "\n".join(item.content for item in references)
    assert "［ 4 ］" in "\n".join(item.content for item in references)
    assert all(
        item.metadata["graph_extraction_enabled"] is False
        for item in references
    )
    assert biographies
    assert all(
        item.metadata["graph_extraction_enabled"] is False
        for item in biographies
    )
    assert any(
        "本文总结系统设计" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_paper_isolates_next_article_preview_interleaved_with_old_references(tmp_path):
    source = tmp_path / "journal-export.md"
    source.write_text(
        "# 第 1 页\n"
        "【摘要】主文章介绍一种协作方法、验证结果、适用限制和后续研究方向。\n"
        "【关键词】协作；验证；方法\n"
        "【Abstract】The main article presents a collaborative method, verified results, "
        "applicable limits, and future research directions.\n"
        "【Keywords】collaboration; verification; method\n"
        "# 正文\n" + "主文章正文详细说明架构、协议、实验、结果与限制。" * 38 + "\n"
        "# 参考文献\n"
        "[1] Zhang Q. Main article reference[J]. Journal, 2024, 10(1): 1-8.\n"
        "[2] Li R. Another reference[J]. Journal, 2025, 11(2): 9-16.\n"
        "# 第 8 页\n"
        "下一篇文章标题\n"
        "【摘要】下一篇文章介绍另一个完全不同的系统。\n"
        "【关键词】另一系统；独立主题\n"
        "Doi: 10.1234/example.next\n"
        "【Abstract】The next article introduces a completely different system.\n"
        "【Keywords】another system; separate topic\n"
        "[3] Wang T. Continued old reference[J]. Journal, 2023, 12(3): 17-24.\n"
        "[4] Zhao M. Final old reference[J]. Journal, 2022, 13(4): 25-32.\n",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )
    preview = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "adjacent_article_preview"
    ]

    assert preview
    assert all(
        item.metadata["graph_extraction_enabled"] is False for item in preview
    )
    assert preview[0].metadata["graph_exclusion_reason"] == (
        "adjacent_article_preview_interleaved_with_references"
    )
    assert "下一篇文章标题" in "\n".join(item.content for item in preview)
    assert any(
        "主文章正文" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_paper_isolates_split_bilingual_next_article_preview_before_old_references(tmp_path):
    source = tmp_path / "split-journal-export.md"
    source.write_text(
        "# 主文章\n"
        "【摘要】主文章介绍一种医学协作方法、验证结果、适用限制和后续方向。\n"
        "【关键词】医学协作；验证；方法\n"
        "Doi: 10.1234/example.main\n"
        "# 正文\n" + "主文章正文详细说明架构、协议、实验、结果与限制。" * 38 + "\n"
        "# 参考文献\n"
        "[1] Zhang Q. Main article reference[J]. Journal, 2024, 10(1): 1-8.\n"
        "[2] Li R. Another reference[J]. Journal, 2025, 11(2): 9-16.\n"
        "# 下一篇文章标题\n"
        "【摘要】下一篇文章只在本页预告另一个完全不同的系统。\n"
        "【关键词】另一系统；独立主题\n"
        "Doi: 10.1234/example.next\n"
        "# Next article title\n"
        "【Abstract】The next article previews a completely different system.\n"
        "【Keywords】another system; separate topic\n"
        "[3] Wang T. Continued old reference[J]. Journal, 2023, 12(3): 17-24.\n"
        "[4] Zhao M. Final old reference[J]. Journal, 2022, 13(4): 25-32.\n",
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )
    preview = [
        item for item in result.chunks
        if item.metadata.get("paper_part") == "adjacent_article_preview"
    ]
    preview_text = "\n".join(item.content for item in preview)

    assert preview
    assert "下一篇文章只在本页预告" in preview_text
    assert "The next article previews" in preview_text
    assert "Continued old reference" not in preview_text
    assert all(
        item.metadata["graph_extraction_enabled"] is False for item in preview
    )
    assert any(
        "Continued old reference" in item.content
        and item.metadata.get("paper_part") == "references"
        and item.metadata["graph_extraction_enabled"] is False
        for item in result.chunks
    )
    assert any(
        "主文章正文" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )


def test_paper_does_not_hide_a_complete_second_article_without_interleaved_references(tmp_path):
    source = tmp_path / "multi-article-collection.md"
    source.write_text(
        "# 摘要\n第一篇文章摘要介绍系统设计与验证结果。\n"
        "# 正文\n" + "第一篇文章正文说明方法和结果。" * 25 + "\n"
        "# 参考文献\n"
        "[1] Zhang Q. A reference[J]. Journal, 2024, 10(1): 1-8.\n"
        "[2] Li R. Another reference[J]. Journal, 2025, 11(2): 9-16.\n"
        "# 第二篇文章\n"
        "【摘要】第二篇文章完整介绍另一系统、实验结果和适用限制。\n"
        "【关键词】另一系统；完整文章\n"
        "Doi: 10.1234/example.second\n" +
        "第二篇文章正文继续详细说明架构、方法、实验与结果。" * 25,
        encoding="utf-8",
    )

    result = process_document(
        source, requested_strategy="paper", source_name=source.name,
    )

    assert not any(
        item.metadata.get("paper_part") == "adjacent_article_preview"
        for item in result.chunks
    )
    assert any(
        "第二篇文章正文" in item.content
        and item.metadata["graph_extraction_enabled"] is True
        for item in result.chunks
    )
