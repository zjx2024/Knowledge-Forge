from io import BytesIO

from docx import Document
from PIL import Image, ImageDraw

from knowledge_forge.chunking import build_processing_result, parse_document_structure
from knowledge_forge.config import Config
from knowledge_forge.media import DocumentMediaService


class FakeVisionOllama:
    def __init__(self):
        self.calls = []

    def json_chat(self, **kwargs):
        self.calls.append(kwargs)
        return {
            "title": "系统架构图",
            "description": "图片展示服务 A 向服务 B 发送请求。",
            "ocr_text": "服务 A  服务 B",
            "facts": ["服务 A 和服务 B 出现在系统架构图中"],
            "relationships": ["服务 A --调用--> 服务 B"],
            "knowledge_relevance": "substantive",
            "relevance_reason": "架构图表达了服务之间的调用关系",
        }


def image_bytes(width=640, height=360):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 80, 220, 180), outline="black", width=4)
    draw.rectangle((400, 80, 580, 180), outline="black", width=4)
    draw.line((220, 130, 400, 130), fill="black", width=4)
    buffer = BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def configure_media(monkeypatch, tmp_path):
    monkeypatch.setattr(Config, "files_root", tmp_path)
    monkeypatch.setattr(Config, "media_root", tmp_path / "media")
    monkeypatch.setattr(Config, "http_base_url", "http://files.test")
    monkeypatch.setattr(Config, "image_analysis_enabled", True)
    monkeypatch.setattr(Config, "image_ocr_enabled", False)
    monkeypatch.setattr(Config, "image_min_width", 96)
    monkeypatch.setattr(Config, "image_min_height", 96)
    monkeypatch.setattr(Config, "image_min_pixels", 16000)


def test_docx_image_becomes_a_provenance_rich_knowledge_chunk(monkeypatch, tmp_path):
    configure_media(monkeypatch, tmp_path)
    source = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("系统组成", level=1)
    document.add_paragraph("系统由多个服务组成。")
    document.add_picture(BytesIO(image_bytes()))
    document.save(source)

    plan, parsed = parse_document_structure(source)
    ollama = FakeVisionOllama()
    enriched = DocumentMediaService(ollama).enrich(
        source,
        parsed,
        {"id": 1, "llm_model": "vision-model", "model_max_tokens": 32768},
        tmp_path / "media" / "project_1" / "hash",
    )
    result = build_processing_result(source, plan, enriched, source_name=source.name, chunk_max_tokens=512)

    image_blocks = [block for block in enriched.blocks if block.block_type == "image"]
    assert len(image_blocks) == 1
    assert image_blocks[0].section_title == "系统组成"
    assert "服务 A --调用--> 服务 B" in image_blocks[0].text
    assert image_blocks[0].metadata["url"].startswith("http://files.test/media/")
    assert any("系统架构图" in chunk.content for chunk in result.chunks)
    assert any(chunk.metadata.get("media_assets") for chunk in result.chunks)
    assert ollama.calls[0]["images"]
    assert "axis names, legends, colours, marker shapes" in (
        ollama.calls[0]["system"]
    )


def test_image_analysis_cache_is_scoped_by_model(monkeypatch, tmp_path):
    configure_media(monkeypatch, tmp_path)
    source = tmp_path / "cached.docx"
    document = Document()
    document.add_paragraph("正文")
    document.add_picture(BytesIO(image_bytes()))
    document.save(source)
    _, parsed = parse_document_structure(source)
    ollama = FakeVisionOllama()
    service = DocumentMediaService(ollama)
    media_dir = tmp_path / "media" / "project_1" / "hash"
    project = {"id": 1, "llm_model": "vision-model", "model_max_tokens": 32768}

    service.enrich(source, parsed, project, media_dir)
    service.enrich(source, parsed, project, media_dir)

    assert len(ollama.calls) == 1
    assert list(media_dir.glob("*.analysis.json"))


def test_scanned_pdf_page_is_available_for_visual_knowledge(monkeypatch, tmp_path):
    configure_media(monkeypatch, tmp_path)
    source = tmp_path / "scan.pdf"
    image = Image.open(BytesIO(image_bytes(900, 1200))).convert("RGB")
    image.save(source, "PDF", resolution=120.0)

    plan, parsed = parse_document_structure(source, requested_parser="plaintext_pdf")
    enriched = DocumentMediaService(FakeVisionOllama()).enrich(
        source,
        parsed,
        {"id": 1, "llm_model": "vision-model", "model_max_tokens": 32768},
        tmp_path / "media" / "project_1" / "pdf",
    )

    image_blocks = [block for block in enriched.blocks if block.block_type == "image"]
    assert image_blocks
    assert image_blocks[0].page == 1
    assert enriched.quality["image_count"] >= 1


def test_tiny_decorative_docx_image_is_filtered(monkeypatch, tmp_path):
    configure_media(monkeypatch, tmp_path)
    source = tmp_path / "decorative.docx"
    document = Document()
    document.add_paragraph("正文")
    document.add_picture(BytesIO(image_bytes(20, 20)))
    document.save(source)
    _, parsed = parse_document_structure(source)
    ollama = FakeVisionOllama()

    enriched = DocumentMediaService(ollama).enrich(
        source,
        parsed,
        {"id": 1, "llm_model": "vision-model", "model_max_tokens": 32768},
        tmp_path / "media" / "project_1" / "tiny",
    )

    assert not [block for block in enriched.blocks if block.block_type == "image"]
    assert ollama.calls == []


def test_context_only_image_remains_retrievable_but_not_graph_eligible(
    monkeypatch, tmp_path,
):
    configure_media(monkeypatch, tmp_path)
    source = tmp_path / "context-image.docx"
    document = Document()
    document.add_heading("作者信息", level=1)
    document.add_paragraph("正文仍应参与图谱抽取。")
    document.add_picture(BytesIO(image_bytes()))
    document.save(source)

    class ContextOnlyVision(FakeVisionOllama):
        def json_chat(self, **kwargs):
            result = super().json_chat(**kwargs)
            return {
                **result,
                "title": "人物照片",
                "description": "一张人物肖像照片。",
                "facts": ["人物位于纯色背景前"],
                "relationships": ["人物 --位于--> 纯色背景"],
                "knowledge_relevance": "context_only",
                "relevance_reason": "图片仅用于作者身份展示",
            }

    plan, parsed = parse_document_structure(source)
    enriched = DocumentMediaService(ContextOnlyVision()).enrich(
        source,
        parsed,
        {"id": 1, "llm_model": "vision-model", "model_max_tokens": 32768},
        tmp_path / "media" / "project_1" / "context",
    )
    result = build_processing_result(
        source, plan, enriched, source_name=source.name, chunk_max_tokens=512,
    )

    image_block = next(
        block for block in enriched.blocks if block.block_type == "image"
    )
    assert image_block.metadata["graph_extraction_enabled"] is False
    image_chunk = next(
        chunk for chunk in result.chunks if "人物肖像照片" in chunk.content
    )
    assert image_chunk.metadata["graph_extraction_enabled"] is False
    assert any(
        "正文仍应参与图谱抽取" in chunk.content
        and chunk.metadata["graph_extraction_enabled"] is True
        for chunk in result.chunks
    )
